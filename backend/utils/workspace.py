"""Durable examination workspace, independent of web-worker memory."""
import hashlib
import json
import uuid
import logging
import random
import threading
import time
import requests
from pymongo.errors import PyMongoError
from datetime import datetime, timezone, timedelta
from pathlib import Path
from contextlib import contextmanager

from pymongo import ReturnDocument
from utils.result_validation import validate_result, ReviewRequired

VERSION = 'grading-v3'
ACTIVE = ['queued', 'processing']
MAX_RETRIES = 5
LEASE_SECONDS = 90
logger = logging.getLogger(__name__)
# Bounded striped locks coalesce shared document work inside the coordinator.
CACHE_LOCKS = [threading.RLock() for _ in range(128)]


def now():
    # Match PyMongo's default UTC-naive decoding while avoiding local-time ambiguity.
    return datetime.now(timezone.utc).replace(tzinfo=None)


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def file_hash(path):
    h = hashlib.sha256()
    with open(path, 'rb') as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def safe_path(root, relative):
    root = Path(root).resolve()
    relative = relative.replace('\\', '/')
    if not relative or ':' in relative or relative.startswith('/') or '..' in relative.split('/'):
        raise ValueError('Invalid upload path')
    target = (root / relative).resolve()
    if not target.is_relative_to(root):
        raise ValueError('Upload path is outside its workspace')
    return target


class LeaseLost(TimeoutError):
    """Infrastructure interruption, not an explicit user cancellation."""


class Workspace:
    def __init__(self, db):
        self.db = db
        self.last_served = {}

    def indexes(self):
        self.db.jobs.create_index([('state', 1), ('created_at', 1)])
        self.db.attempts.create_index([('sheet_id', 1), ('created_at', -1)])
        self.db.sheets.create_index([('exam_id', 1), ('course_code', 1)])
        self.db.sheets.create_index([('course_code', 1), ('created_at', -1), ('_id', 1)])
        self.db.workspace_cache.create_index('sheet_id')
        self.db.workspace_cache.create_index('sheet_ids')
        self.db.upload_parts.create_index('session_id')
        self.db.preparations.create_index([('state', 1), ('created_at', 1)])
        self.db.jobs.create_index([('state', 1), ('next_run_at', 1), ('created_at', 1)])
        self.db.jobs.create_index([('state', 1), ('sheet_snapshot.course_code', 1), ('created_at', 1)])
        # Register known courses without overwriting names or granting permissions.
        from utils.course_processor import COURSE_CODE_PATTERN
        for code in set(self.db.sheets.distinct('course_code') + self.db.evaluations.distinct('course_code')):
            if isinstance(code, str) and len(code) <= 80 and COURSE_CODE_PATTERN.fullmatch(code):
                self.db.courses.update_one({'_id': code}, {'$setOnInsert': {
                    'name': code, 'uploaders': [], 'reviewers': [], 'created_at': now()}}, upsert=True)


    def audit(self, actor, action, **details):
        self.db.audit.insert_one({'actor': actor, 'action': action, 'at': now(), **details})

    def prepare(self, directory, exam_id, exam_type, max_marks, actor, issues=None):
        from utils.course_processor import CourseProcessor
        from utils.gemini_service import VLLM_MODEL, END_SEM_PROMPT, MST_PROMPT
        if not isinstance(exam_id, str) or not exam_id.strip() or len(exam_id) > 120:
            raise ValueError('An examination ID/session is required (maximum 120 characters)')
        if exam_type not in ('mst', 'end_sem') or not 0 < max_marks <= 1000:
            raise ValueError('Invalid examination type or maximum marks')
        processor = CourseProcessor(directory, max_marks)
        courses, _ = processor.discover_courses(allow_empty=True)
        issues = issues if issues is not None else []
        issues.extend(processor.discovery_errors)
        for code in processor.discovered_codes:
            self.db.courses.update_one({'_id': code}, {'$setOnInsert': {
                'name': code, 'uploaders': [], 'reviewers': [], 'created_at': now()}}, upsert=True)
        ids = []
        for code, course in courses.items():
            try:
                for student in course['student_sheets']:
                    sid = digest([exam_id.strip(), code, student['roll_number']])
                    if self.db.jobs.find_one({'_id': sid, 'state': {'$in': ACTIVE + ['managing']}}):
                        raise ValueError('Wait for running sheets to finish before uploading corrections')
                    file_hash(student['file_path'])
                key_hash = file_hash(course['model_answer_path'])
                paper_path = course.get('question_paper_path')
                config = {'version': VERSION, 'model': VLLM_MODEL, 'key': key_hash,
                          'exam_type': exam_type, 'max_marks': max_marks,
                          'prompt': digest(MST_PROMPT if exam_type == 'mst' else END_SEM_PROMPT)}
                config['question_paper'] = file_hash(paper_path) if paper_path else None
                for student in course['student_sheets']:
                    sheet_hash = file_hash(student['file_path'])
                    sheet_config = {**config, 'content_hash': sheet_hash}
                    sheet_id = digest([exam_id.strip(), code, student['roll_number']])
                    running = self.db.jobs.find_one({'_id': sheet_id, 'state': {'$in': ACTIVE + ['managing']}})
                    if running:
                        raise ValueError('Wait for this sheet to finish before replacing its uploaded inputs')
                    self.db.sheets.update_one({'_id': sheet_id}, {'$setOnInsert': {
                        'exam_id': exam_id.strip(), 'course_code': code, 'roll_number': student['roll_number'],
                        'active_attempt': None, 'created_at': now(), 'config': sheet_config,
                        'file_path': student['file_path'], 'model_path': course['model_answer_path'],
                        'question_path': paper_path, 'max_marks': max_marks, 'exam_type': exam_type}}, upsert=True)
                    with self.reserve([sheet_id]):
                        previous = self.db.sheets.find_one({'_id': sheet_id})
                        pin = bool(previous and previous.get('active_pinned') and previous['config'] == sheet_config)
                        self.db.sheets.update_one({'_id': sheet_id}, {'$set': {'file_path': student['file_path'], 'model_path': course['model_answer_path'],
                                   'question_path': paper_path,
                                   'content_hash': sheet_hash, 'config': sheet_config, 'max_marks': max_marks,
                                   'exam_type': exam_type,
                                   'active_pinned': pin,
                                   'updated_at': now()}}, upsert=True)
                    ids.append(sheet_id)
            except (ValueError, OSError) as exc:
                issues.append(f'{code}: {exc}')
        self.audit(actor, 'prepare', exam_id=exam_id, count=len(ids))
        return ids

    def enqueue(self, sheet_id, actor, force=False, request_id=None):
        sheet = self.db.sheets.find_one({'_id': sheet_id})
        if not sheet:
            raise ValueError('Sheet not found')
        # One stable job document per sheet provides an atomic duplicate-work guard.
        existing = self.db.jobs.find_one({'_id': sheet_id})
        if existing and existing['state'] in ACTIVE + ['managing']:
            return existing
        active = self.db.attempts.find_one({'_id': sheet.get('active_attempt')})
        if active and active.get('config') == sheet['config'] and active.get('status') == 'completed' and not force:
            return {'_id': sheet_id, 'state': 'reused'}
        attempt_id = digest([request_id, sheet_id]) if request_id else uuid.uuid4().hex
        if request_id and self.db.attempts.find_one({'_id': attempt_id}):
            return {'_id': sheet_id, 'state': 'reused'}
        job = {'state': 'queued', 'stage': 'queued', 'attempt_id': attempt_id,
               'actor': actor, 'force': force, 'cancel_requested': False,
               'sheet_snapshot': sheet,
               'created_at': now(), 'updated_at': now(), 'error': None, 'progress': None, 'retry_count': 0, 'next_run_at': now(),
               'events': [{'stage': 'queued', 'at': now()}]}
        if existing:
            changed = self.db.jobs.update_one({'_id': sheet_id, 'state': existing['state'],
                                               'attempt_id': existing['attempt_id']}, {'$set': job})
            if not changed.modified_count:
                return self.db.jobs.find_one({'_id': sheet_id})
        else:
            from pymongo.errors import DuplicateKeyError
            try:
                self.db.jobs.insert_one({'_id': sheet_id, **job})
            except DuplicateKeyError:
                return self.db.jobs.find_one({'_id': sheet_id})
        self.audit(actor, 'reevaluate' if force else 'evaluate', sheet_id=sheet_id, attempt_id=attempt_id)
        return {'_id': sheet_id, **job}

    def snapshot(self, query=None, *, limit=None, offset=0, summaries=False, include_jobs=True):
        cursor = self.db.sheets.find(query or {}).sort([('created_at', -1), ('_id', 1)]).skip(offset)
        if limit is not None:
            cursor = cursor.limit(limit)
        sheets = list(cursor)
        ids = [s['_id'] for s in sheets]
        attempt_query = {'_id': {'$in': [s.get('active_attempt') for s in sheets]}} if summaries else {'sheet_id': {'$in': ids}}
        attempts_by_sheet = {}
        for attempt in self.db.attempts.find(attempt_query, {'ocr': 0, 'file_path': 0, 'model_path': 0}).sort('created_at', -1):
            attempts_by_sheet.setdefault(attempt['sheet_id'], []).append(attempt)
        jobs = {j['_id']: j for j in self.db.jobs.find({'_id': {'$in': ids}}, {'sheet_snapshot': 0})} if include_jobs else {}
        output = []
        for sheet in sheets:
            attempts = attempts_by_sheet.get(sheet['_id'], [])
            active = next((a for a in attempts if a['_id'] == sheet.get('active_attempt')), None)
            if active and (active.get('config') != sheet['config'] or active.get('status') != 'completed'):
                active = None
            if summaries and active:
                active = {k: v for k, v in active.items() if k in ('_id', 'marks', 'percentage', 'grade', 'verified', 'verified_by')}
            output.append({**{k: v for k, v in sheet.items() if k not in ('file_path', 'model_path', 'question_path', 'execution')},
                           'attempts': [] if summaries else attempts, 'active': active, 'job': jobs.get(sheet['_id'])})
        return output

    def renew(self, owner):
        self.db.sheets.update_many({'execution.owner': owner, 'execution.state': 'processing',
                                   'execution.expires': {'$gt': now()}},
                                  {'$set': {'execution.expires': now() + timedelta(seconds=LEASE_SECONDS)}})

    def page(self, query, page=1, size=50, view='all', results_only=False):
        pipeline = [{'$match': query},
            {'$lookup': {'from': 'attempts', 'localField': 'active_attempt', 'foreignField': '_id', 'as': '_active'}},
            {'$unwind': {'path': '$_active', 'preserveNullAndEmptyArrays': True}},
            {'$lookup': {'from': 'jobs', 'localField': '_id', 'foreignField': '_id', 'as': '_job'}},
            {'$unwind': {'path': '$_job', 'preserveNullAndEmptyArrays': True}},
            {'$project': {'_job.events': 0, '_job.sheet_snapshot': 0, '_active.question_scores': 0, '_active.feedback': 0, '_active.rubric': 0, '_active.strengths': 0, '_active.missing_points': 0}},
            {'$addFields': {'_valid': {'$and': [{'$eq': ['$_active.status', 'completed']}, {'$eq': ['$_active.config', '$config']}]}}}]
        if results_only and view in ('all', 'needs_review'):
            pipeline.append({'$match': {'$or': [{'_valid': True}, {'_job.state': 'needs_review'}]}})
        elif results_only or view in ('existing', 'pending', 'verified'):
            pipeline.append({'$match': {'_valid': True}})
        elif view == 'ungraded':
            pipeline.append({'$match': {'_valid': {'$ne': True}}})
        if view == 'pending':
            pipeline.append({'$match': {'_active.verified': {'$ne': True}}})
        elif view == 'verified':
            pipeline.append({'$match': {'_active.verified': True}})
        elif view == 'failed':
            pipeline.append({'$match': {'_job.state': {'$in': ['error', 'needs_review']}}})
        elif view == 'needs_review':
            pipeline.append({'$match': {'_job.state': 'needs_review'}})
        if results_only and view in ('pending', 'verified'):
            pipeline.append({'$match': {'_job.state': {'$ne': 'needs_review'}}})
        pipeline.extend([{'$addFields': {'_state': {'$cond': [
            {'$in': ['$_job.state', ['needs_review'] if results_only else ['queued', 'processing', 'error', 'cancelled', 'needs_review']]}, '$_job.state',
            {'$cond': ['$_valid', {'$cond': ['$_active.verified', 'verified', 'review']}, 'ready']}]}}},
            {'$facet': {
                'rows': [{'$sort': {'created_at': -1, '_id': 1}}, {'$skip': (page - 1) * size}, {'$limit': size}, {'$project': {'_id': 1}}],
                'counts': [{'$group': {'_id': '$_state', 'count': {'$sum': 1}}}],
                'total': [{'$count': 'value'}]}}])
        data = next(iter(self.db.sheets.aggregate(pipeline)), {})
        ids = [row['_id'] for row in data.get('rows', [])]
        rows = self.snapshot({'_id': {'$in': ids}}, summaries=True)
        return {'sheets': rows, 'page': page, 'page_size': size,
                'total': next(iter(data.get('total', [])), {}).get('value', 0),
                'counts': {row['_id']: row['count'] for row in data.get('counts', [])}}

    def publish(self, job, sheet, owner):
        # The sheet is the authority for publication. Recovery replaces this token
        # before requeueing, and cancellation writes the same document.
        lease = self.db.worker_leases.find_one({'_id': 'inference'})
        if lease and (lease.get('owner') != owner or lease['expires'] <= now()):
            raise InterruptedError('Worker ownership changed before publication')
        condition = {'_id': sheet['_id'], 'config': sheet['config'],
                     'execution.token': job['token'], 'execution.owner': owner,
                     'execution.state': 'processing', 'execution.cancel_requested': False,
                     'execution.expires': {'$gt': now()}}
        # A pinned result remains active; both branches consume the same execution token.
        changes = {'execution.state': 'completed', 'updated_at': now()}
        saved = self.db.sheets.update_one({**condition, 'active_pinned': {'$ne': True}},
                   {'$set': {**changes, 'active_attempt': job['attempt_id']}})
        if not saved.matched_count:
            saved = self.db.sheets.update_one({**condition, 'active_pinned': True}, {'$set': changes})
        if not saved.matched_count:
            raise InterruptedError('Result was superseded, cancelled, or its lease expired')

    @contextmanager
    def reserve(self, ids):
        # Reserve every selected job before deleting anything. No running worker may own one.
        reserved = []
        reservation_owner = uuid.uuid4().hex
        renewal_stop = threading.Event()
        def renew_reservations():
            while not renewal_stop.wait(30):
                try:
                    self.db.jobs.update_many({'state': 'managing', 'reservation_owner': reservation_owner},
                        {'$set': {'reservation_expires': now() + timedelta(minutes=5)}})
                except PyMongoError:
                    logger.exception('Reservation heartbeat failed')
        renewal = threading.Thread(target=renew_reservations, daemon=True)
        renewal.start()
        try:
            for sheet_id in sorted(set(ids)):
                if not self.db.sheets.find_one({'_id': sheet_id}):
                    raise ValueError('Sheet not found')
                job = self.db.jobs.find_one({'_id': sheet_id})
                if job and job['state'] in ACTIVE + ['managing']:
                    raise ValueError('Cancel and wait for active work before managing these sheets')
                if job:
                    result = self.db.jobs.update_one({'_id': sheet_id, 'state': job['state'], 'attempt_id': job['attempt_id']}, {'$set': {
                        'state': 'managing', 'reservation_owner': reservation_owner, 'reservation_expires': now() + timedelta(minutes=5)}})
                    if not result.modified_count:
                        raise ValueError('Sheet changed; refresh and try again')
                else:
                    self.db.jobs.insert_one({'_id': sheet_id, 'state': 'managing', 'attempt_id': '',
                        'reservation_owner': reservation_owner, 'reservation_expires': now() + timedelta(minutes=5)})
                reserved.append((sheet_id, job))
            yield reserved
        finally:
            renewal_stop.set()
            renewal.join(timeout=1)
            for sheet_id, job in reserved:
                if job:
                    self.db.jobs.replace_one({'_id': sheet_id, 'state': 'managing', 'reservation_owner': reservation_owner}, job)
                else:
                    self.db.jobs.delete_one({'_id': sheet_id, 'state': 'managing', 'reservation_owner': reservation_owner})

    def manage(self, ids, action, actor):
        if action not in ('cache', 'records', 'both', 'sheet') or not ids:
            raise ValueError('Choose sheets and cache, records, or both')
        with self.reserve(ids) as reservations:
            if action in ('records', 'both'):
                for _, previous_job in reservations:
                    if previous_job:
                        previous_job.update(state='ready', stage='ready', error=None)
            for sheet_id in set(ids):
                if action in ('cache', 'both', 'sheet'):
                    self.db.sheets.update_one({'_id': sheet_id}, {'$set': {'cache_generation': uuid.uuid4().hex}})
                    self.db.workspace_cache.update_many({'sheet_ids': sheet_id}, {'$pull': {'sheet_ids': sheet_id}})
                    self.db.workspace_cache.delete_many({'$or': [{'sheet_ids': []}, {'sheet_id': sheet_id, 'sheet_ids': {'$exists': False}}]})
                if action in ('records', 'both', 'sheet'):
                    self.db.attempts.delete_many({'sheet_id': sheet_id})
                    self.db.sheets.update_one({'_id': sheet_id}, {'$set': {'active_attempt': None, 'active_pinned': False}})
                if action == 'sheet':
                    self.db.sheets.delete_one({'_id': sheet_id})
                    self.db.jobs.delete_one({'_id': sheet_id, 'state': 'managing'})
            self.audit(actor, 'manage', sheet_ids=ids, scope=action)

    def claim(self, owner):
        self.db.jobs.update_many({'state': 'managing', 'reservation_expires': {'$lt': now()}}, {'$set': {
            'state': 'error', 'stage': 'error', 'error': 'Management operation interrupted; inspect records before retrying'}})
        eligible = {'state': 'queued', '$or': [
            {'next_run_at': {'$exists': False}}, {'next_run_at': {'$lte': now()}}]}
        courses = self.db.jobs.distinct('sheet_snapshot.course_code', eligible)
        if courses:
            course = min(courses, key=lambda code: (self.last_served.get(code, 0), code))
            eligible['sheet_snapshot.course_code'] = course
            self.last_served[course] = time.monotonic()
        job = self.db.jobs.find_one_and_update(eligible, {'$set': {
            'state': 'processing', 'stage': 'starting', 'error': None, 'owner': owner, 'token': uuid.uuid4().hex, 'claimed_at': now(), 'updated_at': now()},
            '$push': {'events': {'$each': [{'stage': 'starting', 'at': now()}], '$slice': -100}}},
            sort=[('created_at', 1)], return_document=ReturnDocument.AFTER)
        if job:
            changed = self.db.sheets.update_one({'_id': job['_id'], 'config': job['sheet_snapshot']['config']}, {'$set': {
                'execution': {'token': job['token'], 'owner': owner, 'state': 'processing',
                              'cancel_requested': False, 'expires': now() + timedelta(seconds=LEASE_SECONDS)}}})
            if not changed.matched_count:
                self.db.jobs.update_one({'_id': job['_id'], 'token': job['token']}, {'$set': {
                    'state': 'error', 'stage': 'error', 'error': 'Inputs changed; prepare and evaluate again'}})
                return None
            self.db.attempts.update_one({'_id': job['attempt_id'], 'status': {'$ne': 'completed'}},
                {'$set': {'status': 'processing', 'token': job['token']}})
        return job

    def claim_preparation(self, owner):
        return self.db.preparations.find_one_and_update({'state': 'queued'}, {'$set': {
            'state': 'processing', 'owner': owner, 'updated_at': now()}},
            sort=[('created_at', 1)], return_document=ReturnDocument.AFTER)

    def process_preparation(self, job):
        try:
            if job.get('kind') == 'evaluate':
                errors = job.get('errors', [])
                ids = job['sheet_ids']
                for index in range(job.get('submitted', 0), len(ids)):
                    try:
                        self.enqueue(ids[index], job['actor'], job.get('force', False), request_id=job['_id'])
                    except ValueError as exc:
                        if len(errors) < 50:
                            errors.append({'sheet_id': ids[index], 'error': str(exc)[:200]})
                    self.db.preparations.update_one({'_id': job['_id'], 'owner': job['owner']},
                        {'$set': {'submitted': index + 1, 'errors': errors, 'updated_at': now()}})
                self.db.preparations.update_one({'_id': job['_id'], 'owner': job['owner']}, {'$set': {
                    'state': 'completed', 'error': 'Some sheets could not be queued; inspect errors' if errors else None, 'updated_at': now()}})
                return
            issues = []
            ids = self.prepare(job['directory'], job['exam_id'], job['exam_type'], job['max_marks'], job['actor'], issues=issues)
            self.db.preparations.update_one({'_id': job['_id'], 'owner': job['owner']}, {'$set': {
                'state': 'completed', 'sheet_ids': ids, 'issues': issues, 'updated_at': now(), 'error': 'Some courses were skipped; upload corrected folders separately' if issues else None}})
        except PyMongoError:
            raise
        except Exception as exc:
            self.db.preparations.update_one({'_id': job['_id'], 'owner': job['owner']}, {'$set': {
                'state': 'error', 'error': str(exc)[:500], 'updated_at': now()}})

    def recover_interrupted(self, owner):
        # Repair only the known old lease-loss cancellation; never revive an
        # explicitly cancelled job or a source/marking error.
        self.db.jobs.update_many({'state': 'cancelled', 'cancel_requested': {'$ne': True},
            'error': 'Worker lease expired; result will not be activated'}, {'$set': {
                'state': 'queued', 'stage': 'retry_wait', 'next_run_at': now(), 'updated_at': now(),
                'error': 'Connection interrupted; resuming saved work automatically'}})
        self.db.preparations.update_many({'state': 'processing', 'owner': {'$ne': owner}},
            {'$set': {'state': 'queued', 'updated_at': now()}})
        for job in self.db.jobs.find({'state': 'processing', 'owner': {'$ne': owner}}):
            sheet = self.db.sheets.find_one({'_id': job['_id']})
            if not sheet:
                continue
            # Fence the old process before changing any attempt or job state.
            token = uuid.uuid4().hex
            self.db.sheets.update_one({'_id': sheet['_id']}, {'$set': {'execution': {
                'token': token, 'owner': owner, 'state': 'processing', 'cancel_requested': bool(job.get('cancel_requested')),
                'expires': now() + timedelta(seconds=LEASE_SECONDS)}}})
            attempt = self.db.attempts.find_one({'_id': job['attempt_id']})
            retries = job.get('retry_count', 0) + 1
            if job.get('cancel_requested'):
                state, message = 'cancelled', 'Evaluation cancelled'
            elif attempt and attempt.get('status') == 'completed':
                self.publish({**job, 'token': token}, sheet, owner)
                state, message = 'completed', None
            else:
                state, message = 'queued', 'Worker interrupted; resuming saved work automatically'
                self.db.attempts.update_one({'_id': job['attempt_id']}, {'$set': {'status': 'processing'}})
            self.db.jobs.update_one({'_id': job['_id'], 'attempt_id': job['attempt_id'],
                'owner': job.get('owner'), 'state': 'processing'}, {'$set': {
                    'state': state, 'stage': state, 'error': message, 'progress': None, 'updated_at': now(),
                    'retry_count': retries, 'next_run_at': now()},
                '$push': {'events': {'$each': [{'stage': state, 'at': now(), 'message': message}], '$slice': -100}}})
            if state != 'completed':
                self.db.sheets.update_one({'_id': sheet['_id'], 'execution.token': token}, {'$set': {'execution.state': state}})
            if state in ('error', 'cancelled'):
                self.db.attempts.update_one({'_id': job['attempt_id'], 'status': {'$ne': 'completed'}},
                    {'$set': {'status': state, 'marks': None, 'percentage': None, 'grade': None,
                              'error': {'type': 'WORKER_INTERRUPTED', 'message': message}, 'finished_at': now()}})

    def process(self, job, pdf, llm, owner):
        sheet = job.get('sheet_snapshot') or self.db.sheets.find_one({'_id': job['_id']})
        started = time.monotonic()
        metrics = {'result_cache_hit': False, 'cache_hits': 0, 'cache_misses': 0, 'extraction_seconds': 0.0, 'grading_seconds': 0.0}
        metrics_lock = threading.Lock()
        def performance():
            return {**metrics, 'processing_seconds': round(time.monotonic() - started, 3)}

        attempt_id = job['attempt_id']
        attempt = {'_id': attempt_id, 'sheet_id': sheet['_id'], 'status': 'processing',
                   'created_at': now(), 'marks': None, 'percentage': None, 'grade': None,
                   'actor': job['actor'], 'token': job['token'], 'config': sheet['config'], 'file_path': sheet['file_path'],
                   'model_path': sheet['model_path']}
        self.db.attempts.update_one({'_id': attempt_id}, {'$setOnInsert': attempt}, upsert=True)

        def check_cancelled():
            if now() - job.get('claimed_at', now()) > timedelta(hours=24):
                raise ReviewRequired('This attempt exceeded its 24-hour processing window; inspect and retry')
            lease = self.db.worker_leases.find_one({'_id': 'inference'})
            if lease and (lease['owner'] != owner or lease['expires'] <= now()):
                raise LeaseLost('Worker connection/lease interrupted; saved work will resume automatically')
            current = self.db.jobs.find_one({'_id': sheet['_id']})
            if (not current or current.get('owner') != owner or current.get('cancel_requested')
                    or current.get('state') != 'processing' or current.get('attempt_id') != attempt_id):
                raise InterruptedError('Evaluation cancelled')

        llm.check_cancelled = check_cancelled

        def stage(name, progress=None):
            check_cancelled()
            self.db.jobs.update_one({'_id': sheet['_id'], 'owner': owner, 'token': job['token'], 'attempt_id': attempt_id, 'state': 'processing'},
                {'$set': {'stage': name, 'updated_at': now(), 'progress': progress},
                 '$push': {'events': {'$each': [{'stage': name, 'at': now()}], '$slice': -100}}})
            logger.info('attempt=%s sheet=%s stage=%s', attempt_id, sheet['_id'], name)

        def cache_read(key):
            value = self.db.workspace_cache.find_one({'_id': key})
            with metrics_lock:
                metrics['cache_hits' if value else 'cache_misses'] += 1
            if value:
                self.db.workspace_cache.update_one({'_id': key}, {'$addToSet': {'sheet_ids': sheet['_id']}})
            return value

        def cache_write(key, values):
            self.db.workspace_cache.update_one({'_id': key}, {'$set': values,
                '$addToSet': {'sheet_ids': sheet['_id']}}, upsert=True)

        def extract(path, kind):
            identity = digest([sheet['exam_id'], sheet['course_code'], kind, path])
            with CACHE_LOCKS[int(identity[:8], 16) % len(CACHE_LOCKS)]:
                return extract_locked(path, kind)

        def extract_locked(path, kind):
            cache_scope = sheet['_id'] if kind == 'student' else digest([sheet['exam_id'], sheet['course_code'], kind])
            key = digest([cache_scope, sheet.get('cache_generation'), kind, file_hash(path), sheet['config']['model'], 'ocr-v4-150dpi'])
            cached = cache_read(key)
            if cached:
                return cached['text']
            stage('extracting_' + kind)
            try:
                text = pdf.extract_text_from_pdf(path)
            except Exception:
                pages = []
                # Real processor renders one missing page at a time. The legacy
                # adapter branch preserves compatibility with third-party processors.
                streaming = hasattr(type(pdf), 'page_count') and hasattr(type(pdf), 'render_page')
                images = None if streaming else pdf.convert_pdf_to_images(path)
                total = pdf.page_count(path) if streaming else len(images)
                try:
                    for index in range(total):
                        stage(f'extracting_{kind}_page_{index + 1}_of_{total}',
                              {'completed': index, 'total': total, 'unit': 'pages'})
                        page_key = digest([key, index])
                        page_cache = cache_read(page_key)
                        if page_cache:
                            page_text = page_cache['text']
                        else:
                            image = pdf.render_page(path, index) if streaming else images[index]
                            try:
                                page_text = pdf.extract_text_from_images_via_gemini([image], llm, document_kind=kind)
                            finally:
                                if streaming:
                                    image.close()
                            if '[UNREADABLE]' in page_text:
                                raise ReviewRequired(f'Unreadable answer text on page {index + 1}; review the scan')
                            if not page_text.strip() or '[Error' in page_text or '[No text detected]' in page_text:
                                raise ValueError(f'Incomplete OCR on page {index + 1}; review the source')
                            cache_write(page_key, {'text': page_text})
                        pages.append(f'Page {index + 1}:\n{page_text}')
                        stage(f'extracted_{kind}_page_{index + 1}_of_{total}',
                              {'completed': index + 1, 'total': total, 'unit': 'pages'})
                    text = '\n'.join(pages)
                finally:
                    for image in images or []:
                        image.close()
            if not text.strip() or '[Error' in text:
                raise ValueError('Incomplete OCR; evaluation stopped')
            cache_write(key, {'text': text})
            return text

        try:
            stage('checking_cache')
            from utils.gemini_service import VLLM_MODEL, END_SEM_PROMPT, MST_PROMPT
            current_prompt = digest(MST_PROMPT if sheet['exam_type'] == 'mst' else END_SEM_PROMPT)
            if (sheet['config']['version'] != VERSION or sheet['config']['model'] != VLLM_MODEL
                    or sheet['config']['prompt'] != current_prompt):
                raise ValueError('Grading configuration changed after preparation; prepare this exam again')
            if (file_hash(sheet['file_path']) != sheet['config']['content_hash']
                    or file_hash(sheet['model_path']) != sheet['config']['key']):
                raise ValueError('Source files changed after preparation; upload and prepare again')
            from utils.marking_schemes import scheme_id
            marking_scheme = self.db.marking_schemes.find_one({'_id': scheme_id(sheet['course_code'], sheet['exam_type'])})
            grade_key = digest([sheet['_id'], 'grade-marking-fallback-v1', sheet.get('cache_generation'), sheet['config'],
                                marking_scheme.get('revision') if marking_scheme else None])
            cached = None if job['force'] else cache_read(grade_key)
            if cached:
                metrics['result_cache_hit'] = True
                result = validate_result(cached['result'], sheet['max_marks'])
            else:
                extraction_started = time.monotonic()
                model_text = extract(sheet['model_path'], 'answer_key')
                student_text = extract(sheet['file_path'], 'student')
                metrics['extraction_seconds'] = round(time.monotonic() - extraction_started, 3)
                stage('grading')
                grading_started = time.monotonic()
                def structured_cache(prompt, value=None):
                    check_cancelled()
                    cache_id = digest(['structured-v1', sheet.get('cache_generation'), sheet['exam_id'], sheet['course_code'], sheet['config']['model'], prompt])
                    if value is None:
                        # Force fresh scores, retaining reusable OCR and mapping.
                        if job.get('force') and prompt.startswith('Grade ONLY question '):
                            return None
                        saved = cache_read(cache_id)
                        return saved.get('json') if saved else None
                    cache_write(cache_id, {'json': value})
                llm.structured_cache = structured_cache
                llm.marking_scheme = marking_scheme if marking_scheme and marking_scheme.get('questions') else None
                def load_question_paper():
                    path = (marking_scheme or {}).get('question_path') or sheet.get('question_path')
                    if not path:
                        # Older prepared sheets did not retain the paper path.
                        # Search only their own course folder, never another course.
                        import re
                        candidates = [p for p in Path(sheet['model_path']).parent.glob('*.pdf')
                                      if re.fullmatch(r'(?:[a-z0-9][a-z0-9_ -]*?[_ -]?)?question[_\s]?paper', p.stem.lower())]
                        if len(candidates) == 1:
                            path = str(candidates[0])
                    if not path or not Path(path).is_file():
                        return None
                    return extract(path, 'question_paper')
                llm.question_paper_loader = load_question_paper
                result = validate_result(llm.evaluate_answer(student_text, model_text, sheet['max_marks'],
                                         question=None, exam_type=sheet['exam_type']), sheet['max_marks'])
                metrics['grading_seconds'] = round(time.monotonic() - grading_started, 3)
                cache_write(grade_key, {'result': result})
            stage('saving')
            saved = self.db.attempts.update_one({'_id': attempt_id, 'status': 'processing', 'token': job['token']}, {'$set': {
                'status': 'completed', 'marks': result['marks_awarded'], 'percentage': result['percentage'],
                'grade': result['grade'], 'feedback': result['feedback'], 'question_scores': result.get('question_scores', []),
                'strengths': result.get('strengths', []), 'missing_points': result.get('missing_points', []),
                'rubric': result.get('rubric'), 'performance': performance(),
                'finished_at': now(), 'error': None}})
            if not saved.matched_count:
                raise InterruptedError('Attempt no longer belongs to this worker')
            self.publish(job, sheet, owner)
            self.db.jobs.update_one({'_id': sheet['_id'], 'owner': owner, 'token': job['token'], 'attempt_id': attempt_id, 'state': 'processing'}, {'$set': {
                'state': 'completed', 'stage': 'completed', 'progress': None, 'updated_at': now()},
                '$push': {'events': {'$each': [{'stage': 'completed', 'at': now()}], '$slice': -100}}})
            logger.info('attempt=%s completed and persisted', attempt_id)
        except Exception as exc:
            persisted = self.db.attempts.find_one({'_id': attempt_id})
            if persisted and persisted.get('status') == 'completed':
                # Keep the valid result and processing job for startup reconciliation.
                raise
            retries = job.get('retry_count', 0) + 1
            infrastructure = isinstance(exc, (requests.exceptions.ConnectionError, requests.exceptions.Timeout, PyMongoError, LeaseLost))
            retryable = isinstance(exc, (requests.exceptions.ConnectionError, requests.exceptions.Timeout, TimeoutError, PyMongoError))
            current_job = self.db.jobs.find_one({'_id': sheet['_id'], 'attempt_id': attempt_id})
            if current_job and current_job.get('cancel_requested'):
                exc = InterruptedError('Evaluation cancelled')
                retryable = False
            status = 'queued' if retryable and (infrastructure or retries <= MAX_RETRIES) else ('cancelled' if isinstance(exc, InterruptedError) else 'needs_review' if isinstance(exc, ReviewRequired) else 'error')
            logger.warning('attempt=%s status=%s error_type=%s', attempt_id, status, type(exc).__name__)
            self.db.attempts.update_one({'_id': attempt_id, 'status': 'processing', 'token': job['token']}, {'$set': {
                'status': 'processing' if status == 'queued' else status, 'error': {'type': type(exc).__name__, 'message': str(exc)[:500]}, 'finished_at': now()}})
            self.db.sheets.update_one({'_id': sheet['_id'], 'execution.token': job['token']}, {'$set': {'execution.state': status}})
            self.db.jobs.update_one({'_id': sheet['_id'], 'owner': owner, 'token': job['token'], 'attempt_id': attempt_id, 'state': 'processing'}, {'$set': {
                'state': status, 'stage': 'retry_wait' if status == 'queued' else status, 'error': str(exc)[:500], 'progress': None, 'updated_at': now(),
                'retry_count': retries, 'next_run_at': now() + timedelta(seconds=min(300, 5 * 2 ** min(retries, 6)) + random.uniform(0, 5))},
                '$push': {'events': {'$each': [{'stage': status, 'at': now(), 'message': str(exc)[:500]}], '$slice': -100}}})


def excel_bytes(rows):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font
    import re
    wb = Workbook()
    wb.remove(wb.active)
    grouped = {}
    for sheet in rows:
        if not (sheet.get('active') or {}).get('verified'):
            continue
        grouped.setdefault(sheet['course_code'], []).append(sheet)
    if not grouped:
        grouped['Results'] = []
    used = set()
    for course, sheets in grouped.items():
        title = re.sub(r'[\\/*?:\[\]]', '_', course)[:25] or 'Course'
        base = title
        index = 1
        while title.lower() in used:
            title = f'{base}_{index}'
            index += 1
        used.add(title.lower())
        ws = wb.create_sheet(title)
        ws.append(['Examination', 'Course', 'Roll number', 'Maximum marks', 'Marks', 'Percentage', 'Grade', 'Status'])
        for sheet in sorted(sheets, key=lambda s: (s['exam_id'], s['roll_number'])):
            result = sheet.get('active') or {}
            ws.append([sheet['exam_id'], course, sheet['roll_number'], sheet['max_marks'], result.get('marks'),
                       result.get('percentage'), result.get('grade'), 'Verified'])
        # Force user-controlled strings to text, including strings beginning with '='.
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column in 'ABCDEFGH':
            ws.column_dimensions[column].width = 22
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

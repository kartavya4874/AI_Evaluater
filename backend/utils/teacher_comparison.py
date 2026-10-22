"""Independent teacher labels and frozen comparisons; never writes AI results."""
import csv
import io
import math
from collections import Counter
from bson import ObjectId
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from zipfile import ZipFile, BadZipFile

COLUMNS = ['source', 'result_id', 'course_code', 'exam_id', 'roll_number', 'question_id', 'teacher_marks', 'max_marks', 'teacher_feedback']
MAX_ROWS = 1000


def text(value):
    return '' if value is None else str(value).strip()


def number(value, label):
    try:
        if isinstance(value, bool):
            raise ValueError()
        result = float(value)
        if not math.isfinite(result):
            raise ValueError()
        return result
    except (TypeError, ValueError):
        raise ValueError(f'{label} must be a finite number')


def parse_upload(upload, course):
    raw = upload.read(5 * 1024 * 1024 + 1)
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError('Upload at most 5 MB')
    name = (upload.filename or '').lower()
    workbook = None
    try:
        if name.endswith('.csv'):
            table = csv.reader(io.StringIO(raw.decode('utf-8-sig')))
        elif name.endswith('.xlsx'):
            with ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError('Expanded workbook exceeds 25 MB')
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            table = workbook.active.iter_rows(values_only=True)
        else:
            raise ValueError('Choose a CSV or .xlsx file')
        headers = [text(v).lower() for v in next(table, [])]
        if len(headers) != len(set(headers)) or not {'teacher_marks', 'max_marks'}.issubset(headers):
            raise ValueError('Use unique headers including teacher_marks and max_marks')
        rows, seen = [], set()
        for line, values in enumerate(table, 2):
            if line > MAX_ROWS + 1:
                raise ValueError(f'Use at most {MAX_ROWS} rows per upload')
            if not any(text(v) for v in values):
                continue
            row = dict(zip(headers, values))
            # Empty teacher cells in the generated template are intentionally skipped.
            if not text(row.get('teacher_marks')):
                continue
            row = {key: text(row.get(key)) for key in COLUMNS}
            if any(len(v) > 2000 for v in row.values()):
                raise ValueError(f'Row {line}: cell is too long')
            if row['course_code'] and row['course_code'] != course:
                raise ValueError(f'Row {line}: course does not match the selected course')
            row['course_code'] = course
            if row['source'] and row['source'] not in ('workspace', 'legacy'):
                raise ValueError(f'Row {line}: source must be workspace or legacy')
            if row['result_id'] and not row['source']:
                raise ValueError(f'Row {line}: result_id requires source')
            if not row['result_id'] and not row['roll_number']:
                raise ValueError(f'Row {line}: provide source/result_id or roll_number')
            row['teacher_marks'] = number(row['teacher_marks'], f'Row {line}: teacher_marks')
            row['max_marks'] = number(row['max_marks'], f'Row {line}: max_marks')
            if not 0 < row['max_marks'] <= 1000 or not 0 <= row['teacher_marks'] <= row['max_marks']:
                raise ValueError(f'Row {line}: marks must be between zero and max_marks (maximum 1000)')
            identity = (row['source'], row['result_id'], row['exam_id'], row['roll_number'], row['question_id'])
            if identity in seen:
                raise ValueError(f'Row {line}: duplicate student/question entry')
            seen.add(identity)
            rows.append({**row, 'row': line})
        if not rows:
            raise ValueError('Enter teacher marks in at least one row')
        return rows
    except (UnicodeError, BadZipFile, csv.Error, InvalidFileException, ParseError, KeyError):
        raise ValueError('Could not read the file; use UTF-8 CSV or a valid .xlsx workbook')
    finally:
        if workbook:
            workbook.close()


def candidates(db, label):
    found = []
    for source, collection in [('workspace', db.sheets), ('legacy', db.evaluations)]:
        if label['source'] and source != label['source']:
            continue
        query = {'course_code': label['course_code']}
        if label['result_id']:
            if source == 'legacy' and not ObjectId.is_valid(label['result_id']):
                continue
            query['_id'] = ObjectId(label['result_id']) if source == 'legacy' else label['result_id']
        else:
            query['$or'] = [{'roll_number': label['roll_number']}, {'student_rollno': label['roll_number']}]
            # Missing examination IDs only match records that also lack one.
            query['exam_id'] = label['exam_id'] if label['exam_id'] else {'$in': [None, '']}
        for record in collection.find(query).limit(2):
            found.append((source, record))
    return found


def compare(db, labels, tolerance):
    output, matched_keys = [], set()
    for label in labels:
        row = {**label, 'status': 'unmatched', 'ai_marks': None, 'difference': None}
        found = candidates(db, label)
        if len(found) != 1:
            row['status'] = 'ambiguous' if found else 'unmatched'
            output.append(row)
            continue
        source, record = found[0]
        row.update(source=source, result_id=str(record['_id']))
        actual_roll = text(record.get('roll_number') or record.get('student_rollno'))
        if (label['roll_number'] and label['roll_number'] != actual_roll) or (label['exam_id'] and label['exam_id'] != text(record.get('exam_id'))):
            row['status'] = 'identity_mismatch'
            output.append(row)
            continue
        row['roll_number'] = actual_roll
        identity = (source, row['result_id'], label['question_id'])
        if identity in matched_keys:
            raise ValueError('Multiple rows resolve to the same result/question; remove duplicates')
        matched_keys.add(identity)
        result = db.attempts.find_one({'_id': record.get('active_attempt'), 'sheet_id': record['_id'],
            'config': record.get('config'), 'status': 'completed'}) if source == 'workspace' else record
        if not result or result.get('status') in ('error', 'processing', 'cancelled') or str(result.get('feedback', '')).startswith(('Error during evaluation:', 'Evaluation error')):
            row['status'] = 'no_ai_result'
        elif result.get('source') == 'manual' or result.get('status') == 'completed_manually':
            row['status'] = 'manual_result'
        else:
            row['attempt_id'] = str(result['_id'])
            row['model'] = (result.get('config') or record.get('config') or {}).get('model', 'Not recorded')
            row['pipeline_version'] = (result.get('config') or record.get('config') or {}).get('version', 'Not recorded')
            score = result
            maximum = record.get('max_marks')
            if label['question_id']:
                scores = [q for q in (result.get('question_scores') or []) if str(q.get('question_id')) == label['question_id']]
                score = scores[0] if len(scores) == 1 else {}
                maximum = score.get('max_marks')
            try:
                maximum = number(maximum, 'AI maximum')
                mark = number(score.get('marks_awarded') if label['question_id'] else score.get('marks'), 'AI marks')
                if not math.isclose(maximum, label['max_marks'], abs_tol=1e-6):
                    row['status'] = 'max_marks_mismatch'
                elif not 0 <= mark <= maximum:
                    row['status'] = 'invalid_ai_marks'
                else:
                    difference = mark - label['teacher_marks']
                    row.update(status='matched', ai_marks=mark, difference=round(difference, 6),
                        difference_pct=round(100 * difference / maximum, 6))
            except ValueError:
                row['status'] = 'missing_question_score' if label['question_id'] else 'invalid_ai_marks'
        output.append(row)
    def metrics(rows):
        matched = [row for row in rows if row['status'] == 'matched']
        count = len(matched)
        relative = [r for r in matched if r['teacher_marks'] > 0]
        return {'uploaded': len(rows), 'compared': count, 'excluded': len(rows) - count,
            'relative_compared': len(relative), 'zero_teacher_rows': count - len(relative),
            'mape_pct': round(sum(100 * abs(r['difference']) / r['teacher_marks'] for r in relative) / len(relative), 3) if relative else None,
            'mae_pct': round(sum(abs(r['difference_pct']) for r in matched) / count, 3) if count else None,
            'bias_pct': round(sum(r['difference_pct'] for r in matched) / count, 3) if count else None,
            'within_tolerance_pct': round(100 * sum(abs(r['difference_pct']) <= tolerance for r in matched) / count, 2) if count else None}
    return {'rows': output, 'totals': metrics([r for r in output if not r['question_id']]),
        'questions': metrics([r for r in output if r['question_id']]), 'statuses': dict(Counter(r['status'] for r in output))}


def compare_runs(previous, current):
    """Measure change only on identical teacher targets with scores in both runs."""
    def identity(row):
        # Original file row stays stable even when an unresolved match later resolves.
        return row['row']

    summaries = {}
    for level in ('totals', 'questions'):
        is_question = level == 'questions'
        old = {identity(r): r for r in previous['rows'] if bool(r['question_id']) == is_question and r['status'] == 'matched'}
        new = {identity(r): r for r in current['rows'] if bool(r['question_id']) == is_question and r['status'] == 'matched'}
        common = set(old) & set(new)
        pairs = [(old[k], new[k]) for k in common if
                 (old[k]['source'], old[k]['result_id'], old[k]['question_id'], old[k]['teacher_marks'], old[k]['max_marks']) ==
                 (new[k]['source'], new[k]['result_id'], new[k]['question_id'], new[k]['teacher_marks'], new[k]['max_marks'])]
        count = len(pairs)
        before = sum(abs(a['difference_pct']) for a, _ in pairs) / count if count else None
        after = sum(abs(b['difference_pct']) for _, b in pairs) / count if count else None
        relative_pairs = [(a, b) for a, b in pairs if a['teacher_marks'] > 0]
        relative_count = len(relative_pairs)
        before_relative = sum(100 * abs(a['difference']) / a['teacher_marks'] for a, _ in relative_pairs) / relative_count if relative_count else None
        after_relative = sum(100 * abs(b['difference']) / b['teacher_marks'] for _, b in relative_pairs) / relative_count if relative_count else None
        summaries[level] = {'paired': count, 'previous_compared': len(old), 'current_compared': len(new),
            'relative_paired': relative_count, 'zero_teacher_pairs': count - relative_count,
            'before_mape_pct': round(before_relative, 3) if relative_count else None,
            'after_mape_pct': round(after_relative, 3) if relative_count else None,
            'mape_change_pct': round(after_relative - before_relative, 3) if relative_count else None,
            'no_longer_comparable': len(old) - count, 'newly_comparable': len(new) - count,
            'before_mae_pct': round(before, 3) if count else None,
            'after_mae_pct': round(after, 3) if count else None,
            'mae_change_pct': round(after - before, 3) if count else None,
            'improved': sum(abs(b['difference_pct']) < abs(a['difference_pct']) - 1e-6 for a, b in pairs),
            'worsened': sum(abs(b['difference_pct']) > abs(a['difference_pct']) + 1e-6 for a, b in pairs),
            'unchanged': sum(math.isclose(abs(a['difference_pct']), abs(b['difference_pct']), abs_tol=1e-6) for a, b in pairs),
            'same_prediction': sum(a.get('attempt_id') == b.get('attempt_id') and a['ai_marks'] == b['ai_marks'] for a, b in pairs)}
    return {'baseline_id': previous['_id'], 'baseline_name': previous['name'],
        'baseline_created_at': previous['created_at'], **summaries}

"""HTTP contract matrix against the registered Flask app, with isolated storage/model doubles."""
import io
import json
import tempfile
import unittest
from pathlib import Path
from datetime import timedelta
from unittest.mock import Mock, patch
import jwt
import mongomock
from bson import ObjectId
import app as module
from utils.workspace import now


class RouteContractTests(unittest.TestCase):
    def setUp(self):
        self.db = mongomock.MongoClient().db
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        for target, value in [('app.db_connection.get_db', self.db)]:
            patcher = patch(target, return_value=value)
            patcher.start(); self.addCleanup(patcher.stop)
        no_network = patch('requests.sessions.Session.request', side_effect=AssertionError('Live network forbidden in API tests'))
        no_network.start(); self.addCleanup(no_network.stop)
        config = patch.dict(module.app.config, {'UPLOAD_FOLDER': self.temp.name, 'CACHE_FOLDER': self.temp.name})
        config.start(); self.addCleanup(config.stop)
        self.client = module.app.test_client()
        self.admin = self.headers(module.Config.ADMIN_EMAIL)
        self.faculty_email = 'faculty@example.com'
        self.db.faculty.insert_one({'_id': self.faculty_email, 'email': self.faculty_email, 'role': 'faculty', 'active': True})
        self.faculty = self.headers(self.faculty_email)
        self.oid = str(ObjectId())
        self.db.teachers.insert_one({'_id': ObjectId(self.oid), 'name': 'Teacher', 'email': 't@example.com'})
        self.db.students.insert_one({'_id': ObjectId(self.oid), 'name': 'Student', 'email': 's@example.com', 'roll_number': '201'})
        self.result = {'_id': ObjectId(self.oid), 'status': 'completed', 'marks': 5, 'max_marks': 10,
            'percentage': 50, 'feedback': 'Checked', 'student_id': self.oid, 'teacher_id': self.oid,
            'course_code': 'MB3204', 'batch_mode': True, 'created_at': now()}
        self.db.evaluations.insert_one(dict(self.result))
        self.db.courses.insert_one({'_id': 'MB3204', 'uploaders': [self.faculty_email], 'reviewers': [self.faculty_email]})
        self.db.sheets.insert_one({'_id': 'sheet', 'course_code': 'MB3204', 'exam_id': 'exam', 'roll_number': '201', 'config': {}, 'max_marks': 10})
        self.db.preparations.insert_one({'_id': 'operation', 'actor': self.faculty_email, 'state': 'completed', 'created_at': now()})

    def headers(self, email):
        return {'Authorization': 'Bearer ' + jwt.encode({'email': email, 'provider': 'google'}, module.Config.JWT_SECRET_KEY, algorithm='HS256')}

    def concrete(self, rule):
        path = rule.rule
        for name in rule.arguments:
            replacement = self.oid if name in ('teacher_id', 'student_id', 'evaluation_id') else {'course_code': 'MB3204', 'sheet_id': 'sheet', 'operation_id': 'operation', 'sid': 'a' * 32, 'index': '0'}.get(name, 'missing')
            import re
            path = re.sub(r'<(?:[^:>]+:)?' + name + r'>', replacement, path)
        return path

    def test_every_private_route_rejects_missing_and_invalid_auth(self):
        public = {'root', 'health_check', 'readiness', 'auth_config', 'auth_login', 'auth_google', 'static'}
        checked = 0
        for rule in module.app.url_map.iter_rules():
            if rule.endpoint in public:
                continue
            for method in rule.methods - {'OPTIONS', 'HEAD'}:
                for headers in ({}, {'Authorization': 'Bearer invalid'}):
                    with self.subTest(path=rule.rule, method=method, headers=bool(headers)):
                        response = self.client.open(self.concrete(rule), method=method, headers=headers)
                        self.assertEqual(response.status_code, 401, response.get_data(as_text=True))
                        self.assertIn('error', response.json)
                checked += 1
        self.assertGreater(checked, 40)

    def test_every_route_preflight_responds_without_credentials(self):
        for rule in module.app.url_map.iter_rules():
            if rule.endpoint == 'static':
                continue
            with self.subTest(path=rule.rule):
                response = self.client.options(self.concrete(rule), headers={'Origin': 'http://localhost:3000', 'Access-Control-Request-Method': next(iter(rule.methods - {'OPTIONS', 'HEAD'}))})
                self.assertEqual(response.status_code, 200)

    def test_all_admin_get_routes_have_expected_response_shapes(self):
        expected = {
            '/': ('message', 200), '/api/auth/config': ('client_id', 200), '/api/auth/me': ('user', 200),
            '/api/health': ('database', 200), '/api/ready': ('worker', 503), '/api/teachers': ('teachers', 200),
            f'/api/teachers/{self.oid}': ('teacher', 200), '/api/students': ('students', 200),
            f'/api/students/{self.oid}': ('student', 200), f'/api/students/{self.oid}/statistics': ('statistics', 200),
            '/api/evaluations': (None, 200), f'/api/evaluations/{self.oid}': ('evaluation', 200),
            f'/api/evaluations/student/{self.oid}': ('evaluations', 200), f'/api/evaluations/teacher/{self.oid}': ('evaluations', 200),
            '/api/evaluations/recent': ('evaluations', 200), '/api/batch/stream': ('error', 410), '/api/batch/status': ('error', 410),
            '/api/batch/results': ('courses', 200), '/api/batch/results/MB3204': ('evaluations', 200),
            '/api/workspace/sheets': ('sheets', 200), '/api/workspace/sheets/sheet': ('sheet', 200),
            '/api/workspace/preparations/operation': ('operation', 200), '/api/access': ('courses', 200),
        }
        for path, (key, status) in expected.items():
            with self.subTest(path=path):
                response = self.client.get(path, headers=self.admin)
                self.assertEqual(response.status_code, status, response.get_data(as_text=True))
                if key:
                    self.assertIn(key, response.json)
                else:
                    self.assertIsInstance(response.json, list)
        response = self.client.get('/api/workspace/export?source=workspace', headers=self.admin)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data.startswith(b'PK'))
        self.assertIn('attachment', response.headers['Content-Disposition'])

    def test_json_endpoints_reject_non_object_and_malformed_bodies(self):
        paths = ['/api/teachers', '/api/students', f'/api/evaluations/{self.oid}/manual', '/api/workspace/prepare',
            '/api/workspace/evaluate', '/api/workspace/evaluate-all', '/api/workspace/cancel', '/api/workspace/manage',
            '/api/workspace/result', '/api/workspace/verify', '/api/workspace/uploads', '/api/access/settings',
            '/api/access/faculty', '/api/access/courses']
        for path in paths:
            for body in ('[]', 'null', '"text"', '1', '{bad'):
                with self.subTest(path=path, body=body):
                    response = self.client.open(path, method='PUT' if path.endswith('/manual') else 'POST', headers=self.admin, data=body, content_type='application/json')
                    self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
                    self.assertIn('error', response.json)

    def test_invalid_ids_and_query_limits_are_client_errors(self):
        paths = ['/api/teachers/bad', '/api/students/bad', '/api/evaluations/bad', '/api/students/bad/statistics', '/api/evaluations/student/bad', '/api/evaluations/teacher/bad']
        for path in paths:
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path, headers=self.admin).status_code, 400)
        for value in ('0', '-1', '1001', 'bad', '1.5'):
            self.assertEqual(self.client.get('/api/evaluations/recent?limit=' + value, headers=self.admin).status_code, 400)
        for query in ('page=0', 'page_size=101', 'offset=-1'):
            path = '/api/workspace/sheets/sheet' if query.startswith('offset') else '/api/workspace/sheets'
            self.assertEqual(self.client.get(path + '?' + query, headers=self.admin).status_code, 400)

    def test_legacy_crud_and_statistics_roundtrip(self):
        for entity in ('teachers', 'students'):
            response = self.client.post('/api/' + entity, headers=self.admin, json={'name': 'New', 'email': 'NEW@example.com'})
            self.assertEqual(response.status_code, 201)
            item = response.json[entity[:-1]]
            self.assertEqual(item['email'], 'new@example.com')
            self.assertEqual(self.client.get('/api/' + entity + '/' + item['_id'], headers=self.admin).status_code, 200)
            self.assertEqual(self.client.delete('/api/' + entity + '/' + item['_id'], headers=self.admin).status_code, 200)
            self.assertEqual(self.client.delete('/api/' + entity + '/' + item['_id'], headers=self.admin).status_code, 404)
        self.db.evaluations.insert_one({'student_id': self.oid, 'status': 'error', 'marks': 0, 'max_marks': 10})
        stats = self.client.get(f'/api/students/{self.oid}/statistics', headers=self.admin).json['statistics']
        self.assertEqual(stats['average_marks'], 5)
        self.assertEqual(stats['total_evaluations'], 1)
        response = self.client.put(f'/api/evaluations/{self.oid}/manual', headers=self.admin, json={'marks_awarded': 7, 'feedback': 'Teacher correction'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.db.evaluations.find_one({'_id': ObjectId(self.oid)})['marks'], 7)
        self.assertEqual(self.client.delete(f'/api/evaluations/{self.oid}', headers=self.admin).status_code, 200)
        self.assertEqual(self.client.delete(f'/api/evaluations/{self.oid}', headers=self.admin).status_code, 404)

    def test_invalid_fields_cannot_become_database_query_operators(self):
        for path, data in [('/api/teachers', {'name': 'x', 'email': {'$ne': None}}),
            ('/api/students', {'name': ['x'], 'email': 'x@example.com'}),
            ('/api/workspace/prepare', {'directory': [], 'exam_id': 'exam', 'exam_type': 'mst', 'max_marks': 10}),
            ('/api/workspace/evaluate', {'sheet_ids': ['sheet'], 'force': 'false'}),
            (f'/api/evaluations/{self.oid}/manual', {'marks_awarded': True})]:
            with self.subTest(path=path):
                response = self.client.open(path, method='PUT' if path.endswith('/manual') else 'POST', headers=self.admin, json=data)
                self.assertEqual(response.status_code, 400)

    def test_scanned_upload_and_single_evaluation_use_ocr_and_clean_sources(self):
        image = Mock()
        result = {'marks_awarded': 5, 'feedback': 'Checked', 'strengths': [], 'missing_points': []}
        with patch.object(module.pdf_processor, 'extract_text_from_pdf', side_effect=ValueError('scan')), patch.object(module.pdf_processor, 'page_count', return_value=1), patch.object(module.pdf_processor, 'render_page', return_value=image), patch.object(module.pdf_processor, 'extract_text_from_images_via_gemini', return_value='Q1 10 marks complete answer') as ocr, patch.object(module.gemini_service, 'evaluate_answer', return_value=result):
            response = self.client.post('/api/upload-model-answer', headers=self.admin, data={'file': (io.BytesIO(b'key-pdf'), 'key.pdf')})
            self.assertEqual(response.status_code, 200, response.json)
            self.assertEqual(ocr.call_args.kwargs['document_kind'], 'answer_key')
            response = self.client.post('/api/evaluate-answer', headers=self.admin, data={'student_file': (io.BytesIO(b'student-pdf'), 'student.pdf'), 'model_answer': 'Q1 10 marks answer', 'max_marks': '10', 'teacher_id': self.oid, 'student_id': self.oid})
            self.assertEqual(response.status_code, 200, response.json)
            self.assertEqual(response.json['evaluation']['marks_awarded'], 5)
            self.assertEqual(ocr.call_args.kwargs['document_kind'], 'student')
            response = self.client.post('/api/ocr-only', headers=self.admin, data={'file': (io.BytesIO(b'ocr-pdf'), 'ocr.pdf')})
            self.assertEqual(response.status_code, 200)
        self.assertEqual(list(Path(self.temp.name).glob('*.pdf')), [])
        self.assertEqual(image.close.call_count, 3)

    def test_upload_failure_cleans_temporary_file_and_does_not_save_grade(self):
        with patch.object(module.pdf_processor, 'extract_text_from_pdf', side_effect=ValueError('scan')), patch.object(module.pdf_processor, 'page_count', return_value=1), patch.object(module.pdf_processor, 'render_page', side_effect=ValueError('bad PDF')):
            response = self.client.post('/api/upload-model-answer', headers=self.admin, data={'file': (io.BytesIO(b'bad'), 'key.pdf')})
            self.assertEqual(response.status_code, 400)
        self.assertEqual(list(Path(self.temp.name).glob('*.pdf')), [])
        self.assertEqual(self.db.evaluations.count_documents({}), 1)

    def test_every_admin_only_route_rejects_faculty(self):
        public = {'root', 'health_check', 'readiness', 'auth_config', 'auth_login', 'auth_google', 'static', 'auth_me'}
        for rule in module.app.url_map.iter_rules():
            if rule.endpoint in public or rule.rule.startswith(('/api/workspace/', '/api/access')) or rule.rule == '/api/batch/upload-folder':
                continue
            for method in rule.methods - {'OPTIONS', 'HEAD'}:
                with self.subTest(path=rule.rule, method=method):
                    self.assertEqual(self.client.open(self.concrete(rule), method=method, headers=self.faculty).status_code, 403)

    def test_upload_prepare_evaluate_review_verify_export_delete_request_chain(self):
        import hashlib
        from flask import Flask
        from workspace_api import create_workspace_api
        from upload_api import create_upload_api
        from werkzeug.exceptions import HTTPException
        isolated = Flask(__name__)
        isolated.register_blueprint(create_workspace_api(module.require_auth, self.temp.name))
        isolated.register_blueprint(create_upload_api(module.require_auth, self.temp.name))
        isolated.register_error_handler(HTTPException, module.http_error)
        client = isolated.test_client()
        base = '/api/workspace'
        sid = 'a' * 32
        manifest = {'session_id': sid, 'files': [{'path': 'MB3204/model_answer.pdf', 'size': 3}, {'path': 'MB3204/student_201.pdf', 'size': 3}]}
        response = client.post(base + '/uploads', headers=self.faculty, json=manifest)
        self.assertEqual(response.status_code, 200, response.json)
        for index in range(2):
            response = client.put(f'{base}/uploads/{sid}/{index}', headers={**self.faculty, 'X-Content-SHA256': hashlib.sha256(b'pdf').hexdigest()}, data=b'pdf')
            self.assertEqual(response.status_code, 200, response.json)
        response = client.post(f'{base}/uploads/{sid}/complete', headers=self.faculty)
        self.assertEqual(response.status_code, 200, response.json)
        payload = {'directory': response.json['upload_dir'], 'exam_id': 'new-exam', 'exam_type': 'mst', 'max_marks': 10}
        response = client.post(base + '/prepare', headers=self.faculty, json=payload)
        self.assertEqual(response.status_code, 200, response.json)
        sheet_id = response.json['sheet_ids'][0]
        response = client.post(base + '/prepare', headers=self.faculty, json={**payload, 'async': True})
        self.assertEqual(response.status_code, 202)
        operation_id = response.json['operation_id']
        self.assertEqual(client.get(base + '/preparations/' + operation_id, headers=self.faculty).status_code, 200)
        response = client.post(base + '/evaluate', headers=self.faculty, json={'sheet_ids': [sheet_id], 'force': False})
        self.assertEqual(response.status_code, 202, response.json)
        self.assertEqual(client.post(base + '/cancel', headers=self.faculty, json={'sheet_ids': [sheet_id]}).status_code, 200)
        self.assertEqual(client.post(base + '/result', headers=self.faculty, json={'sheet_id': sheet_id, 'marks': 5, 'reason': 'Teacher review'}).status_code, 200)
        detail = client.get(base + '/sheets/' + sheet_id, headers=self.faculty).json['sheet']
        attempt_id = detail['active']['_id']
        self.assertEqual(client.post(base + '/result', headers=self.faculty, json={'sheet_id': sheet_id, 'attempt_id': attempt_id}).status_code, 200)
        response = client.post(base + '/verify', headers=self.faculty, json={'source': 'workspace', 'items': [{'id': sheet_id, 'attempt_id': attempt_id}]})
        self.assertEqual(response.status_code, 200, response.json)
        response = client.get(base + '/export?source=workspace', headers=self.faculty)
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(response.data))
        self.assertGreater(workbook.active.max_row, 1)
        workbook.close()
        self.assertEqual(client.post(base + '/evaluate-all', headers=self.faculty, json={'exam_id': 'new-exam'}).status_code, 202)
        self.assertEqual(client.post(base + '/manage', headers=self.admin, json={'sheet_ids': [sheet_id], 'scope': 'sheet', 'confirmation': 'DELETE'}).status_code, 200)
        self.assertIsNone(self.db.sheets.find_one({'_id': sheet_id}))
        self.assertTrue((Path(self.temp.name) / sid / 'MB3204/student_201.pdf').exists())

    def test_admin_assignment_import_and_legacy_verification_requests(self):
        response = self.client.post('/api/access/faculty', headers=self.admin, json={'email': 'new@example.com', 'name': 'New faculty', 'role': 'faculty', 'active': True})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.post('/api/access/courses', headers=self.admin, json={'code': 'MB3204', 'name': 'Course', 'uploaders': ['new@example.com'], 'reviewers': ['new@example.com']}).status_code, 200)
        self.assertEqual(self.client.post('/api/access/settings', headers=self.admin, json={'parallelism': 2}).status_code, 200)
        from openpyxl import Workbook
        workbook = Workbook(); workbook.active.append(['course_code', 'name']); workbook.active.append(['MB3205', 'Next course'])
        content = io.BytesIO(); workbook.save(content); workbook.close(); content.seek(0)
        self.assertEqual(self.client.post('/api/access/import', headers=self.admin, data={'file': (content, 'courses.xlsx')}).status_code, 200)
        from utils.verification import revision
        response = self.client.post('/api/workspace/verify', headers=self.admin, json={'source': 'legacy', 'items': [{'id': self.oid, 'revision': revision(self.result)}]})
        self.assertEqual(response.status_code, 200, response.json)
        for path in ('/api/batch/process-courses', '/api/batch/stop', '/api/auth/login'):
            self.assertEqual(self.client.post(path, headers=self.admin, json={}).status_code, 410)

    def test_missing_json_content_type_and_wrong_methods(self):
        for path in ('/api/teachers', '/api/students', '/api/workspace/evaluate', '/api/access/courses'):
            response = self.client.post(path, headers=self.admin, data='{}', content_type='text/plain')
            self.assertEqual(response.status_code, 415)
            self.assertIn('error', response.json)
        self.assertEqual(self.client.patch('/api/workspace/sheets', headers=self.admin).status_code, 405)
        self.assertEqual(self.client.get('/api/no-such-route', headers=self.admin).status_code, 404)

    def test_frontend_literal_api_requests_match_registered_methods_and_paths(self):
        import re
        adapter = module.app.url_map.bind('localhost')
        source_root = Path(__file__).resolve().parents[1] / 'frontend' / 'src'
        checked = 0
        pattern = re.compile(r"api\.(get|post|put|delete|patch)\(\s*(['\"`])(.*?)\2", re.DOTALL)
        for source in source_root.rglob('*'):
            if source.suffix not in ('.js', '.jsx') or '.test.' in source.name:
                continue
            content = source.read_text(encoding='utf-8')
            constants = dict(re.findall(r"const\s+(\w+)\s*=\s*['\"](/[^'\"]+)['\"]", content))
            for method, quote, literal in pattern.findall(content):
                path = '/api' + re.sub(r'\$\{([^}]+)\}', lambda match: constants.get(match.group(1), '0'), literal).split('?')[0]
                with self.subTest(source=source.name, method=method, path=path):
                    adapter.match(path, method=method.upper())
                checked += 1
        self.assertGreater(checked, 40)

    def test_manual_legacy_revision_clears_stale_ai_totals_before_verification(self):
        self.db.evaluations.update_one({'_id': ObjectId(self.oid)}, {'$set': {'question_scores': [{'question_id': '1', 'max_marks': 10, 'marks_awarded': 5}]}})
        response = self.client.put(f'/api/evaluations/{self.oid}/manual', headers=self.admin, json={'marks_awarded': 7, 'feedback': 'Teacher correction'})
        self.assertEqual(response.status_code, 200)
        record = self.db.evaluations.find_one({'_id': ObjectId(self.oid)})
        from utils.verification import revision
        response = self.client.post('/api/workspace/verify', headers=self.admin, json={'source': 'legacy', 'items': [{'id': self.oid, 'revision': revision(record)}]})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(len(self.db.audit.find_one({'action': 'manual_legacy_grade'})['previous_question_scores']), 1)

    def test_storage_outage_returns_json_503_even_during_authentication(self):
        with patch('app.db_connection.get_db', side_effect=RuntimeError('private connection details')):
            for path in ('/api/access', '/api/workspace/sheets', '/api/teachers'):
                response = self.client.get(path, headers=self.admin)
                self.assertEqual(response.status_code, 503)
                self.assertNotIn('private', response.get_data(as_text=True))


if __name__ == '__main__':
    unittest.main()

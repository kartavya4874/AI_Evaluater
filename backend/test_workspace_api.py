import tempfile
import unittest
from functools import wraps
from unittest.mock import patch
import mongomock
from flask import Flask, g, request, jsonify

with patch('pymongo.MongoClient', mongomock.MongoClient):
    from workspace_api import create_workspace_api


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.db = mongomock.MongoClient().db
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.patch = patch('workspace_api.db_connection.get_db', return_value=self.db)
        self.patch.start()
        self.addCleanup(self.patch.stop)
        app = Flask(__name__)
        def auth(fn):
            @wraps(fn)
            def wrapped(*args, **kwargs):
                if request.headers.get('Authorization') != 'Bearer test':
                    return jsonify(error='unauthorized'), 401
                g.user_email = 'admin'
                g.user = {'role': request.headers.get('X-Test-Role', 'admin')}
                return fn(*args, **kwargs)
            return wrapped
        app.register_blueprint(create_workspace_api(auth, self.temp.name))
        self.client = app.test_client()
        self.headers = {'Authorization': 'Bearer test'}
        self.db.courses.insert_one({'_id': 'FS3401'})
        self.db.sheets.insert_one({'_id': 'sheet', 'config': {}, 'max_marks': 100,
                                  'exam_id': 'exam', 'course_code': 'FS3401', 'roll_number': '101'})

    def test_result_lifecycle_stays_consistent_across_review_inventory_and_export(self):
        from io import BytesIO
        from openpyxl import load_workbook
        def get(path):
            response = self.client.get('/api/workspace/' + path, headers=self.headers)
            self.assertEqual(response.status_code, 200)
            return response.get_json()
        def post(path, body):
            response = self.client.post('/api/workspace/' + path, headers=self.headers, json=body)
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        post('result', {'sheet_id': 'sheet', 'marks': 76, 'reason': 'Teacher review'})
        row = get('results')['items'][0]
        self.assertEqual(get('inventory')['items'][0]['marks'], row['result']['marks'])
        post('verify', {'source': 'workspace', 'items': [{'id': 'sheet', 'attempt_id': row['result']['_id']}]})
        self.assertEqual(get('results?view=verified')['total'], 1)
        self.assertTrue(get('inventory')['items'][0]['verified'])
        exported = self.client.get('/api/workspace/export', headers=self.headers)
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(BytesIO(exported.data))
        self.assertTrue(any(76 in row for sheet in workbook for row in sheet.values))
        post('result', {'sheet_id': 'sheet', 'marks': 80, 'reason': 'Corrected grading'})
        self.assertEqual(get('results?view=verified')['total'], 0)
        self.assertFalse(get('inventory')['items'][0]['verified'])
        post('manage', {'sheet_ids': ['sheet'], 'scope': 'records', 'confirmation': 'DELETE'})
        self.assertEqual(get('results')['total'], 0)
        self.assertEqual(get('inventory?status=no_marks')['total'], 1)
        post('manage', {'sheet_ids': ['sheet'], 'scope': 'sheet', 'confirmation': 'DELETE'})
        self.assertEqual(get('inventory')['total'], 0)

    def test_all_routes_require_auth(self):
        for path in ('sheets', 'export', 'inventory', 'results'):
            self.assertEqual(self.client.get('/api/workspace/' + path).status_code, 401)
        for path in ('prepare', 'evaluate', 'cancel', 'manage', 'result', 'verify'):
            self.assertEqual(self.client.post('/api/workspace/' + path, json={}).status_code, 401)

    def test_inventory_includes_both_stores_and_sheets_without_marks(self):
        from bson import ObjectId
        saved_id = ObjectId()
        self.db.evaluations.insert_one({'_id': saved_id, 'marks': 76, 'max_marks': 100,
            'student_rollno': 'single-101', 'batch_mode': False})
        response = self.client.get('/api/workspace/inventory', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data['total'], 2)
        self.assertEqual(data['items'][0]['source'], 'workspace')
        self.assertIsNone(data['items'][0]['marks'])
        self.assertEqual(data['items'][1]['id'], str(saved_id))
        self.assertEqual(data['items'][1]['roll_number'], 'single-101')

    def test_inventory_requires_administrator(self):
        response = self.client.get('/api/workspace/inventory', headers={**self.headers, 'X-Test-Role': 'faculty'})
        self.assertEqual(response.status_code, 403)

    def test_search_matches_both_stores_and_treats_regex_characters_literally(self):
        self.client.post('/api/workspace/result', headers=self.headers, json={'sheet_id': 'sheet', 'marks': 85, 'reason': 'Review'})
        self.db.evaluations.insert_many([
            {'course_code': 'FS3401', 'student_rollno': '101-old', 'marks': 72, 'max_marks': 100},
            {'course_code': 'FS3401', 'roll_number': 'a.b', 'marks': 60, 'max_marks': 100},
            {'course_code': 'FS3401', 'roll_number': 'axb', 'marks': 60, 'max_marks': 100}])
        for endpoint in ('results', 'inventory'):
            matched = self.client.get(f'/api/workspace/{endpoint}?q=101', headers=self.headers).get_json()
            self.assertEqual(matched['total'], 2)
            literal = self.client.get(f'/api/workspace/{endpoint}?q=a.b', headers=self.headers).get_json()
            self.assertEqual(literal['total'], 1)
            source = self.client.get(f'/api/workspace/{endpoint}?q=101&source=legacy', headers=self.headers).get_json()
            self.assertEqual(source['total'], 1)
            self.assertEqual(source['items'][0]['source'], 'legacy')
            invalid = self.client.get(f'/api/workspace/{endpoint}?q=' + 'x' * 121, headers=self.headers)
            self.assertEqual(invalid.status_code, 400)

    def test_inventory_filters_ungraded_and_failed_across_sources(self):
        self.db.evaluations.insert_one({'course_code': 'FS3401', 'roll_number': 'failed-old', 'marks': 0, 'max_marks': 100, 'status': 'error'})
        self.db.jobs.insert_one({'_id': 'sheet', 'state': 'error'})
        failed = self.client.get('/api/workspace/inventory?status=failed', headers=self.headers).get_json()
        self.assertEqual(failed['total'], 2)
        ungraded = self.client.get('/api/workspace/inventory?status=no_marks', headers=self.headers).get_json()
        self.assertEqual(ungraded['total'], 2)
        self.assertTrue(all(row['marks'] is None for row in ungraded['items']))
        exam = self.client.get('/api/workspace/inventory?exam_id=exam', headers=self.headers).get_json()
        self.assertEqual(exam['total'], 1)

    def test_search_export_and_permissions_use_the_same_filter(self):
        import io
        self.db.evaluations.insert_many([{'course_code': code, 'roll_number': roll, 'marks': 70, 'max_marks': 100,
            'verified': True, 'status': 'completed', 'feedback': 'Reviewed'} for code, roll in [('FS3401', '001'), ('OTHER', '001'), ('FS3401', '002')]])
        self.db.courses.update_one({'_id': 'FS3401'}, {'$set': {'reviewers': ['admin']}})
        headers = {**self.headers, 'X-Test-Role': 'faculty'}
        for endpoint in ('results',):
            data = self.client.get(f'/api/workspace/{endpoint}?q=001', headers=headers).get_json()
            self.assertEqual(data['total'], 1)
        with patch('workspace_api.excel_bytes', return_value=io.BytesIO(b'workbook')) as export:
            response = self.client.get('/api/workspace/export?source=legacy&q=001', headers=headers)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(len(export.call_args.args[0]), 1)
            self.assertEqual(export.call_args.args[0][0]['course_code'], 'FS3401')

    def test_combined_results_count_previous_records_when_workspace_is_empty(self):
        self.db.evaluations.insert_many([{'marks': 76, 'max_marks': 100, 'course_code': 'FS3401', 'verified': i < 2} for i in range(10)])
        response = self.client.get('/api/workspace/results', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual((data['total'], data['verified'], data['pending']), (10, 2, 8))
        self.assertEqual(len(data['items']), 10)

    def test_combined_results_paginates_filters_and_preserves_verification_identity(self):
        self.client.post('/api/workspace/result', headers=self.headers, json={'sheet_id': 'sheet', 'marks': 85, 'reason': 'Review'})
        self.db.evaluations.insert_one({'marks': 76, 'max_marks': 100, 'course_code': 'FS3401', 'verified': True})
        first = self.client.get('/api/workspace/results?page_size=1', headers=self.headers).get_json()
        second = self.client.get('/api/workspace/results?page_size=1&page=2', headers=self.headers).get_json()
        self.assertEqual(first['total'], 2)
        self.assertEqual(first['items'][0]['source'], 'workspace')
        self.assertTrue(first['items'][0]['result']['_id'])
        self.assertEqual(second['items'][0]['source'], 'legacy')
        self.assertTrue(second['items'][0]['result']['verification_revision'])
        verified = self.client.get('/api/workspace/results?view=verified', headers=self.headers).get_json()
        self.assertEqual(verified['total'], 1)
        self.assertEqual(verified['items'][0]['source'], 'legacy')
        filtered = self.client.get('/api/workspace/results?course_code=OTHER', headers=self.headers).get_json()
        self.assertEqual(filtered['total'], 0)

    def test_combined_results_scope_and_detail_protect_unassigned_records(self):
        from bson import ObjectId
        record_id = ObjectId()
        self.db.evaluations.insert_one({'_id': record_id, 'marks': 76, 'max_marks': 100, 'course_code': 'FS3401'})
        headers = {**self.headers, 'X-Test-Role': 'faculty'}
        data = self.client.get('/api/workspace/results', headers=headers).get_json()
        self.assertEqual(data['total'], 0)
        self.assertEqual(self.client.get(f'/api/workspace/results/legacy/{record_id}', headers=headers).status_code, 403)
        self.db.courses.update_one({'_id': 'FS3401'}, {'$set': {'reviewers': ['admin']}})
        data = self.client.get('/api/workspace/results', headers=headers).get_json()
        self.assertEqual(data['total'], 1)
        self.assertEqual(self.client.get(f'/api/workspace/results/legacy/{record_id}', headers=headers).status_code, 200)

    def test_clearing_marks_keeps_sheet_visible_for_further_management(self):
        self.client.post('/api/workspace/result', headers=self.headers,
            json={'sheet_id': 'sheet', 'marks': 76, 'reason': 'Reviewed'})
        before = self.client.get('/api/workspace/inventory', headers=self.headers).get_json()
        self.assertEqual(before['items'][0]['marks'], 76)
        response = self.client.post('/api/workspace/manage', headers=self.headers,
            json={'sheet_ids': ['sheet'], 'scope': 'records', 'confirmation': 'DELETE'})
        self.assertEqual(response.status_code, 200)
        after = self.client.get('/api/workspace/inventory', headers=self.headers).get_json()
        self.assertEqual(after['total'], 1)
        self.assertEqual(after['items'][0]['id'], 'sheet')
        self.assertIsNone(after['items'][0]['marks'])

    def test_inventory_pagination_crosses_collections_without_omissions(self):
        self.db.evaluations.insert_one({'marks': 42, 'course_code': 'FS3401'})
        pages = [self.client.get(f'/api/workspace/inventory?page_size=1&page={page}', headers=self.headers).get_json()
                 for page in (1, 2, 3)]
        self.assertEqual([len(page['items']) for page in pages], [1, 1, 0])
        self.assertEqual(pages[0]['items'][0]['source'], 'workspace')
        self.assertEqual(pages[1]['items'][0]['source'], 'legacy')
        only_saved = self.client.get('/api/workspace/inventory?source=legacy', headers=self.headers).get_json()
        self.assertEqual(only_saved['total'], 1)
        missing_course = self.client.get('/api/workspace/inventory?course_code=OTHER', headers=self.headers).get_json()
        self.assertEqual(missing_course['total'], 0)

    def test_manual_revision_persists_consistent_result(self):
        response = self.client.post('/api/workspace/result', headers=self.headers,
            json={'sheet_id': 'sheet', 'marks': 50, 'reason': 'Teacher review'})
        self.assertEqual(response.status_code, 200)
        sheet = self.db.sheets.find_one({'_id': 'sheet'})
        result = self.db.attempts.find_one({'_id': sheet['active_attempt']})
        self.assertEqual(result['marks'], 50)
        self.assertEqual(result['percentage'], 50)
        self.assertTrue(sheet['active_pinned'])

    def test_invalid_manual_marks_rejected(self):
        response = self.client.post('/api/workspace/result', headers=self.headers,
            json={'sheet_id': 'sheet', 'marks': 101, 'reason': 'Teacher review'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.db.attempts.count_documents({}), 0)

    def test_delete_requires_confirmation(self):
        response = self.client.post('/api/workspace/manage', headers=self.headers,
            json={'sheet_ids': ['sheet'], 'scope': 'both'})
        self.assertEqual(response.status_code, 400)

    def test_operator_injection_rejected(self):
        response = self.client.post('/api/workspace/result', headers=self.headers,
            json={'sheet_id': {'$ne': None}, 'marks': 0, 'reason': 'x'})
        self.assertEqual(response.status_code, 400)

    def test_server_directory_outside_uploads_rejected(self):
        response = self.client.post('/api/workspace/prepare', headers=self.headers,
            json={'directory': '..', 'exam_id': 'exam', 'exam_type': 'mst', 'max_marks': 50})
        self.assertEqual(response.status_code, 400)

    def test_snapshot_reports_worker_offline(self):
        response = self.client.get('/api/workspace/sheets', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json['worker_online'])

    def test_verify_legacy_bulk_and_export_only_verified(self):
        from utils.verification import revision
        from openpyxl import load_workbook
        from io import BytesIO
        records = []
        for roll in ('101', '102', '103'):
            r = {'roll_number': roll, 'course_code': 'FS3401', 'marks': 60, 'max_marks': 100,
                 'percentage': 60, 'grade': 'B', 'feedback': 'Full private feedback', 'batch_mode': True}
            r['_id'] = self.db.evaluations.insert_one(r).inserted_id
            records.append(r)
        response = self.client.post('/api/workspace/verify', headers=self.headers, json={
            'source': 'legacy', 'items': [{'id': str(r['_id']), 'revision': revision(r)} for r in records[:2]]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.db.evaluations.count_documents({'verified': True}), 2)
        exported = self.client.get('/api/workspace/export', headers=self.headers)
        values = list(load_workbook(BytesIO(exported.data)).active.values)
        self.assertEqual(len(values), 3)
        self.assertEqual([v[2] for v in values[1:]], ['101', '102'])
        self.assertNotIn('Full private feedback', str(values))

    def test_stale_legacy_verification_rejected(self):
        from utils.verification import revision
        r = {'marks': 60, 'max_marks': 100, 'feedback': 'Reviewed'}
        identifier = self.db.evaluations.insert_one(r).inserted_id
        old = revision(r)
        self.db.evaluations.update_one({'_id': identifier}, {'$set': {'marks': 70}})
        response = self.client.post('/api/workspace/verify', headers=self.headers, json={
            'source': 'legacy', 'items': [{'id': str(identifier), 'revision': old}]})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.db.evaluations.count_documents({'verified': True}), 0)

    def test_failed_grade_cannot_be_verified(self):
        from utils.verification import revision
        record = {'marks': 0, 'max_marks': 100, 'feedback': 'Evaluation error: timeout'}
        identifier = self.db.evaluations.insert_one(record).inserted_id
        response = self.client.post('/api/workspace/verify', headers=self.headers, json={
            'source': 'legacy', 'items': [{'id': str(identifier), 'revision': revision(record)}]})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.db.evaluations.count_documents({'verified': True}), 0)

    def test_workspace_verification_does_not_transfer_to_new_revision(self):
        from openpyxl import load_workbook
        from io import BytesIO
        self.client.post('/api/workspace/result', headers=self.headers,
            json={'sheet_id': 'sheet', 'marks': 50, 'reason': 'Review'})
        attempt = self.db.sheets.find_one({'_id': 'sheet'})['active_attempt']
        response = self.client.post('/api/workspace/verify', headers=self.headers, json={
            'source': 'workspace', 'items': [{'id': 'sheet', 'attempt_id': attempt}]})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(self.db.attempts.find_one({'_id': attempt})['verified'])
        self.client.post('/api/workspace/result', headers=self.headers,
            json={'sheet_id': 'sheet', 'marks': 60, 'reason': 'Changed grade'})
        response = self.client.post('/api/workspace/verify', headers=self.headers, json={
            'source': 'workspace', 'items': [{'id': 'sheet', 'attempt_id': attempt}]})
        self.assertEqual(response.status_code, 400)
        exported = self.client.get('/api/workspace/export', headers=self.headers)
        self.assertEqual(load_workbook(BytesIO(exported.data)).active.max_row, 1)


if __name__ == '__main__':
    unittest.main()

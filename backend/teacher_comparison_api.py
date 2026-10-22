import io
import uuid
from flask import Blueprint, request, jsonify, g, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pymongo.errors import PyMongoError
from werkzeug.exceptions import Forbidden, NotFound
from access import authorize, scope
from utils.db_connection import db_connection
from utils.workspace import now
from utils.teacher_comparison import COLUMNS, MAX_ROWS, parse_upload, compare, compare_runs, number


def create_teacher_comparison_api(require_auth):
    api = Blueprint('teacher_comparisons', __name__, url_prefix='/api/workspace/teacher-comparisons')

    @api.before_request
    def guard():
        if request.method != 'OPTIONS':
            return require_auth(lambda: None)()

    @api.errorhandler(ValueError)
    def invalid(error):
        return jsonify(error=str(error)), 400

    @api.errorhandler(PyMongoError)
    @api.errorhandler(RuntimeError)
    def unavailable(error):
        return jsonify(error='Persistent storage is unavailable. Refresh before retrying.'), 503

    def get_report(report_id):
        db = db_connection.get_db()
        report = db.teacher_comparisons.find_one({'_id': report_id})
        if not report:
            raise NotFound('Comparison not found')
        authorize(db, report['course_code'], 'reviewers')
        if report['state'] == 'draft' and report['actor'] != g.user_email:
            raise Forbidden('This preview belongs to another reviewer')
        return report

    def check_dataset_overlap(db, course, split, rows):
        identities = {(r['source'], r['result_id']) for r in rows if r['result_id'] and r['status'] not in ('unmatched', 'ambiguous', 'identity_mismatch')}
        for previous in db.teacher_comparisons.find({'course_code': course, 'state': 'saved', 'split': {'$ne': split}}, {'rows.source': 1, 'rows.result_id': 1}):
            if identities.intersection((r['source'], r['result_id']) for r in previous.get('rows', [])):
                raise ValueError('These papers already belong to the other dataset purpose. Keep calibration and held-out test papers separate.')

    @api.get('')
    def listing():
        db = db_connection.get_db()
        courses = sorted(set(filter(None, db.sheets.distinct('course_code', scope(db, {}, 'reviewers')) + db.evaluations.distinct('course_code', scope(db, {}, 'reviewers')))))
        reports = list(db.teacher_comparisons.find(scope(db, {'state': 'saved'}, 'reviewers'), {'rows': 0, 'labels': 0}).sort('created_at', -1).limit(100))
        return jsonify(courses=courses, reports=reports)

    def workbook_bytes(headers, rows, summary=None):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Teacher marks'
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for cells in ws:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Never interpret imported notes or identifiers as formulas.
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='243B53')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(45, max(18, len(str(column[0].value)) + 3))
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        if summary:
            overview = wb.create_sheet('Comparison summary', 0)
            overview.append(['Measure', 'Value'])
            for entry in summary:
                overview.append(entry)
            for cells in overview:
                for cell in cells:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
            for cell in overview[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='243B53')
            overview.column_dimensions['A'].width = 58
            overview.column_dimensions['B'].width = 65
            overview.freeze_panes = 'A2'
            wb.active = 0
        result = io.BytesIO()
        wb.save(result)
        result.seek(0)
        return result

    @api.get('/template')
    def template():
        db = db_connection.get_db()
        course = request.args.get('course_code', '').strip()
        if not course:
            raise ValueError('Choose a course')
        authorize(db, course, 'reviewers')
        level = request.args.get('level', 'totals')
        if level not in ('totals', 'questions'):
            raise ValueError('Choose totals or questions')
        rows = []
        for source, collection in [('workspace', db.sheets), ('legacy', db.evaluations)]:
            for record in collection.find({'course_code': course}).sort([('created_at', -1), ('_id', 1)]).limit(MAX_ROWS + 1):
                result = db.attempts.find_one({'_id': record.get('active_attempt'), 'sheet_id': record['_id']}) if source == 'workspace' else record
                scores = ((result or {}).get('question_scores') or []) if level == 'questions' else [{'question_id': '', 'max_marks': record.get('max_marks')}]
                for score in scores:
                    rows.append([source, str(record['_id']), course, record.get('exam_id', ''),
                        str(record.get('roll_number') or record.get('student_rollno') or ''), str(score.get('question_id', '')),
                        None, score.get('max_marks'), ''])
                    if len(rows) > MAX_ROWS:
                        raise ValueError('This course exceeds 1000 template rows. Upload a CSV for a smaller examination instead.')
        if not rows:
            raise ValueError('No saved question scores available; use the totals template first' if level == 'questions' else 'No records found for this course')
        return send_file(workbook_bytes(COLUMNS, rows), as_attachment=True, download_name='teacher-marks-template.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @api.post('/preview')
    def preview():
        db = db_connection.get_db()
        course = request.form.get('course_code', '').strip()
        name = request.form.get('name', '').strip()
        split = request.form.get('split', 'calibration')
        independent = request.form.get('independent', 'false')
        if not course or not name or len(name) > 120 or split not in ('calibration', 'test') or independent not in ('true', 'false'):
            raise ValueError('Enter a comparison name, course and valid dataset purpose')
        authorize(db, course, 'reviewers')
        tolerance = number(request.form.get('tolerance', '3'), 'Tolerance')
        if not 0 <= tolerance <= 100:
            raise ValueError('Tolerance must be between 0 and 100 percentage points')
        upload = request.files.get('file')
        if not upload:
            raise ValueError('Choose a teacher marks CSV or .xlsx file')
        labels = parse_upload(upload, course)
        comparison = compare(db, labels, tolerance)
        check_dataset_overlap(db, course, split, comparison['rows'])
        report = {'_id': uuid.uuid4().hex, 'course_code': course, 'name': name, 'split': split,
            'independent': independent == 'true', 'tolerance': tolerance, 'actor': g.user_email, 'teacher_actor': g.user_email,
            'created_at': now(), 'state': 'draft', 'labels': labels, **comparison}
        db.teacher_comparisons.insert_one(report)
        return jsonify(report={k: v for k, v in report.items() if k != 'labels'})

    @api.get('/<report_id>')
    def detail(report_id):
        return jsonify(report={k: v for k, v in get_report(report_id).items() if k != 'labels'})

    @api.post('/<report_id>/save')
    def save(report_id):
        report = get_report(report_id)
        if report['actor'] != g.user_email:
            raise Forbidden('Only the uploading reviewer may save this preview')
        check_dataset_overlap(db_connection.get_db(), report['course_code'], report['split'], report['rows'])
        db_connection.get_db().teacher_comparisons.update_one({'_id': report_id, 'state': 'draft'}, {'$set': {'state': 'saved', 'saved_at': now()}})
        return jsonify(success=True)

    @api.post('/<report_id>/compare-latest')
    def compare_latest(report_id):
        original = get_report(report_id)
        if original['state'] != 'saved':
            raise ValueError('Save the teacher benchmark before comparing again')
        db = db_connection.get_db()
        report = {k: original[k] for k in ('course_code', 'name', 'split', 'independent', 'tolerance', 'labels', 'teacher_actor')}
        report.update(_id=uuid.uuid4().hex, parent_id=original['_id'], actor=g.user_email,
            created_at=now(), state='draft', **compare(db, original['labels'], original['tolerance']))
        report['change'] = compare_runs(original, report)
        db.teacher_comparisons.insert_one(report)
        return jsonify(report={k: v for k, v in report.items() if k != 'labels'})

    @api.get('/<report_id>/export')
    def export(report_id):
        report = get_report(report_id)
        columns = COLUMNS + ['status', 'ai_marks', 'difference', 'difference_pct', 'attempt_id', 'model', 'pipeline_version']
        summary = [['Comparison', report['name']], ['Course', report['course_code']],
            ['Created (UTC)', str(report['created_at'])], ['Dataset purpose', report['split']],
            ['Teacher independence confirmed', report['independent']], ['Tolerance (percentage points)', report['tolerance']],
            ['Interpretation', 'Differences are AI minus teacher. Excluded rows are not counted as correct.']]
        for level in ('totals', 'questions'):
            for metric, value in report[level].items():
                summary.append([f'{level}: {metric}', value])
        if report.get('change'):
            summary.append(['Baseline comparison', report['change']['baseline_name']])
            summary.append(['Paired change', 'Uses only identical targets compared in both runs; negative MAE change means closer to teacher.'])
            for level in ('totals', 'questions'):
                for metric, value in report['change'][level].items():
                    summary.append([f'Paired {level}: {metric}', value])
        return send_file(workbook_bytes(columns, [[row.get(c) for c in columns] for row in report['rows']], summary),
            as_attachment=True, download_name='teacher-comparison.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    return api

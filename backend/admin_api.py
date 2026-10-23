from utils.api_contract import json_object
from config import Config
from flask import Blueprint, request, jsonify, g
from werkzeug.exceptions import Forbidden
from openpyxl import load_workbook
from utils.db_connection import db_connection
from utils.runtime_settings import parallelism, MAX_PARALLELISM
from utils.workspace import now
from access import is_admin, codes, course_code
from utils.marking_schemes import validate_scheme, scheme_id
from io import BytesIO
from pathlib import Path
import uuid


def create_admin_api(require_auth):
    api = Blueprint('access', __name__, url_prefix='/api/access')

    @api.errorhandler(ValueError)
    def invalid(e):
        return jsonify(error=str(e)), 400

    @require_auth
    def authenticated_guard():
        if request.method != 'GET' and not is_admin():
            raise Forbidden('Administrator access required')

    @api.before_request
    def guard():
        # Browser preflights carry no bearer token; Flask supplies the OPTIONS response.
        if request.method == 'OPTIONS':
            return None
        return authenticated_guard()

    @api.get('')
    def listing():
        db = db_connection.get_db()
        query = {} if is_admin() else {'_id': {'$in': codes(db)}}
        return jsonify(bootstrap_admin=Config.ADMIN_EMAIL.strip().lower() if is_admin() else None, parallelism=parallelism(db), courses=list(db.courses.find(query)),
                       faculty=list(db.faculty.find({}, {'google_sub': 0})) if is_admin() else [])

    @api.get('/marking-scheme')
    def get_marking_scheme():
        if not is_admin():
            raise Forbidden('Administrator access required')
        if request.args.get('exam_type') not in ('mst', 'end_sem'):
            raise ValueError('Choose MST or End Sem')
        identifier = scheme_id(request.args.get('course_code', ''), request.args.get('exam_type', ''))
        row = db_connection.get_db().marking_schemes.find_one({'_id': identifier}) or {}
        row['has_question_paper'] = bool(row.pop('question_path', None))
        return jsonify(scheme=row)

    @api.post('/marking-scheme')
    def save_marking_scheme():
        scheme = validate_scheme(json_object())
        db = db_connection.get_db()
        identifier = scheme_id(scheme['course_code'], scheme['exam_type'])
        scheme.update(updated_by=g.user_email, updated_at=now(), revision=uuid.uuid4().hex)
        db.marking_schemes.update_one({'_id': identifier}, {'$set': scheme}, upsert=True)
        db.audit.insert_one({'actor': g.user_email, 'action': 'save_marking_scheme', 'scheme_id': identifier,
                             'revision': scheme['revision'], 'at': now()})
        return jsonify(success=True)

    @api.post('/marking-scheme/import')
    def preview_marking_import():
        from zipfile import ZipFile, BadZipFile
        upload = request.files.get('file')
        if not upload or not upload.filename.lower().endswith('.xlsx'):
            raise ValueError('Upload an .xlsx workbook')
        raw = upload.read(2 * 1024 * 1024 + 1)
        if len(raw) > 2 * 1024 * 1024:
            raise ValueError('Workbook must be smaller than 2 MB')
        try:
            with ZipFile(BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('Expanded workbook is too large')
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
            try:
                rows = list(workbook.active.iter_rows(min_row=1, max_row=202, max_col=5, values_only=True))
            finally:
                workbook.close()
        except (BadZipFile, KeyError, OSError) as exc:
            raise ValueError('Invalid workbook') from exc
        headers = ('question_id', 'max_marks', 'group', 'choose', 'selection')
        if not rows or tuple(rows[0]) != headers:
            raise ValueError('Columns must be: question_id, max_marks, group, choose, selection')
        questions = []
        for values in rows[1:]:
            if all(v is None for v in values):
                continue
            if any(isinstance(v, str) and v.startswith('=') for v in values):
                raise ValueError('Use literal values, not Excel formulas')
            row = dict(zip(headers, values))
            if row['selection'] is None:
                row['selection'] = 'exact'
            questions.append(row)
        maximum = request.form.get('max_marks', '')
        try:
            maximum = float(maximum)
        except ValueError:
            raise ValueError('Enter maximum marks before importing') from None
        scheme = validate_scheme({'course_code': request.form.get('course_code'),
            'exam_type': request.form.get('exam_type'), 'max_marks': maximum, 'questions': questions})
        return jsonify(scheme=scheme)

    @api.post('/marking-scheme/paper')
    def save_default_paper():
        from PyPDF2 import PdfReader
        code, exam_type = course_code(request.form.get('course_code', '')), request.form.get('exam_type')
        if exam_type not in ('mst', 'end_sem'):
            raise ValueError('Choose MST or End Sem')
        upload = request.files.get('file')
        if not upload or not upload.filename.lower().endswith('.pdf'):
            raise ValueError('Upload a PDF question paper')
        raw = upload.read(20 * 1024 * 1024 + 1)
        if len(raw) > 20 * 1024 * 1024:
            raise ValueError('Question paper must be smaller than 20 MB')
        try:
            if not 1 <= len(PdfReader(BytesIO(raw)).pages) <= 200:
                raise ValueError('Invalid page count')
        except Exception as exc:
            raise ValueError('Upload a readable PDF with 1 to 200 pages') from exc
        root = Path(__file__).resolve().parent / Config.UPLOAD_FOLDER / 'marking_papers'
        root.mkdir(parents=True, exist_ok=True)
        path = root / (uuid.uuid4().hex + '.pdf')
        path.write_bytes(raw)
        db = db_connection.get_db()
        revision = uuid.uuid4().hex
        db.marking_schemes.update_one({'_id': scheme_id(code, exam_type)}, {'$set': {
            'course_code': code, 'exam_type': exam_type, 'question_path': str(path),
            'revision': revision, 'updated_at': now(), 'updated_by': g.user_email}}, upsert=True)
        db.audit.insert_one({'actor': g.user_email, 'action': 'save_default_question_paper',
                            'scheme_id': scheme_id(code, exam_type), 'revision': revision, 'at': now()})
        return jsonify(success=True)

    @api.post('/settings')
    def settings():
        value = (json_object()).get('parallelism')
        if type(value) is not int or not 1 <= value <= MAX_PARALLELISM:
            raise ValueError('Choose between 1 and 5 parallel evaluations')
        db = db_connection.get_db()
        db.settings.update_one({'_id': 'evaluation'}, {'$set': {
            'parallelism': value, 'updated_by': g.user_email, 'updated_at': now()}}, upsert=True)
        return jsonify(success=True, parallelism=value)

    @api.post('/faculty')
    def faculty():
        data = json_object()
        if not isinstance(data.get('email'), str) or not isinstance(data.get('name'), str):
            raise ValueError('Enter a name and valid Google email address')
        email = data['email'].strip().lower()
        name = data['name'].strip()
        if not name or len(name) > 120 or not __import__('re').fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
            raise ValueError('Enter a name and valid Google email address')
        db = db_connection.get_db()
        existing = db.faculty.find_one({'_id': email}) or {}
        role = data.get('role', existing.get('role', 'faculty'))
        active = data.get('active', existing.get('active', True))
        if role not in ('admin', 'faculty') or type(active) is not bool:
            raise ValueError('Choose a valid account role and access status')
        if email == Config.ADMIN_EMAIL.strip().lower():
            raise ValueError('The initial administrator is managed through ADMIN_EMAIL in the server settings')
        if email == g.user_email and (role != 'admin' or not active):
            raise ValueError('You cannot remove your own administrator access')
        db.faculty.update_one({'_id': email}, {'$set': {'email': email, 'name': name,
            'role': role, 'active': active}}, upsert=True)
        db.audit.insert_one({'actor': g.user_email, 'action': 'update_account',
            'email': email, 'role': role, 'active': active, 'at': now()})
        return jsonify(success=True)

    @api.post('/courses')
    def courses():
        db = db_connection.get_db()
        data = json_object()
        code = course_code(data.get('code'))
        assignments = {}
        for field in ('uploaders', 'reviewers'):
            values = data.get(field, [])
            if not isinstance(values, list) or not all(isinstance(v, str) for v in values):
                raise ValueError('Invalid faculty selection')
            values = list(set(v.strip().lower() for v in values))
            if any(not db.faculty.find_one({'_id': v, 'active': True}) for v in values):
                raise ValueError('Assignments must use active faculty accounts')
            assignments[field] = values
        db.courses.update_one({'_id': code}, {'$set': {'name': str(data.get('name', '')).strip()[:200], **assignments}}, upsert=True)
        return jsonify(success=True)

    @api.post('/import')
    def import_courses():
        upload = request.files.get('file')
        if not upload or not upload.filename.lower().endswith('.xlsx'):
            raise ValueError('Choose an .xlsx workbook with course_code and optional name columns')
        try:
            wb = load_workbook(upload, read_only=True, data_only=False)
            rows = wb.active.iter_rows(values_only=True)
            headers = [str(v or '').strip().lower() for v in next(rows)]
            if 'course_code' not in headers:
                raise ValueError('The first row must contain course_code')
            prepared = {}
            for number, row in enumerate(rows, 2):
                if number > 5001:
                    raise ValueError('Import at most 5000 rows')
                value = row[headers.index('course_code')]
                if value is None:
                    continue
                code = course_code(value)
                name = str(row[headers.index('name')] or '')[:200] if 'name' in headers else ''
                prepared[code] = name
            wb.close()
            if not prepared:
                raise ValueError('No course codes found')
        except ValueError:
            raise
        except Exception:
            raise ValueError('Unable to read workbook; use a valid .xlsx file')
        db = db_connection.get_db()
        for code, name in prepared.items():
            db.courses.update_one({'_id': code}, {'$set': {'name': name}, '$setOnInsert': {'uploaders': [], 'reviewers': []}}, upsert=True)
        return jsonify(success=True, count=len(prepared))

    return api

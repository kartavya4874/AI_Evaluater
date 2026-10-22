import io
import unittest
from functools import wraps
from unittest.mock import patch
import mongomock
from bson import ObjectId
from flask import Flask, request, g, jsonify
from openpyxl import load_workbook, Workbook
from werkzeug.datastructures import FileStorage
with patch('pymongo.MongoClient', mongomock.MongoClient):
    from teacher_comparison_api import create_teacher_comparison_api
    from utils.teacher_comparison import parse_upload


class TeacherComparisonTests(unittest.TestCase):
    def setUp(self):
        self.db = mongomock.MongoClient().db
        patcher = patch('teacher_comparison_api.db_connection.get_db', return_value=self.db)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.db.courses.insert_one({'_id': 'FS3401', 'reviewers': ['teacher@example.com']})
        self.db.sheets.insert_one({'_id': 'sheet', 'course_code': 'FS3401', 'exam_id': 'October', 'roll_number': '001', 'max_marks': 100, 'active_attempt': 'attempt', 'config': {'model': 'model-v1', 'version': 'grading-v3'}})
        self.db.attempts.insert_one({'_id': 'attempt', 'sheet_id': 'sheet', 'config': {'model': 'model-v1', 'version': 'grading-v3'}, 'status': 'completed', 'marks': 76, 'question_scores': [{'question_id': '1', 'max_marks': 10, 'marks_awarded': 8}]})
        self.legacy = ObjectId()
        self.db.evaluations.insert_one({'_id': self.legacy, 'course_code': 'FS3401', 'roll_number': '002', 'max_marks': 50, 'marks': 30})
        app = Flask(__name__)
        def auth(fn):
            @wraps(fn)
            def wrapped(*args, **kwargs):
                if request.headers.get('Authorization') != 'Bearer test':
                    return jsonify(error='unauthorized'), 401
                g.user = {'role': 'faculty'}
                g.user_email = request.headers.get('X-User', 'teacher@example.com')
                return fn(*args, **kwargs)
            return wrapped
        app.register_blueprint(create_teacher_comparison_api(auth))
        self.client = app.test_client()
        self.headers = {'Authorization': 'Bearer test'}
        self.base = '/api/workspace/teacher-comparisons'

    def preview(self, csv_text=None, **overrides):
        text = csv_text or f'source,result_id,roll_number,teacher_marks,max_marks,question_id\nworkspace,sheet,001,72,100,\nlegacy,{self.legacy},002,32,50,\nworkspace,sheet,001,7,10,1\n'
        data = dict(course_code='FS3401', name='Teacher benchmark', split='calibration', independent='true', tolerance='3')
        data.update(overrides)
        data['file'] = (io.BytesIO(text.encode()), 'marks.csv')
        return self.client.post(self.base + '/preview', headers=self.headers, data=data)

    def test_metrics_separate_totals_questions_and_preserve_ai(self):
        response = self.preview()
        self.assertEqual(response.status_code, 200, response.get_json())
        report = response.get_json()['report']
        self.assertEqual(report['totals']['compared'], 2)
        self.assertEqual(report['totals']['mae_pct'], 4)
        self.assertEqual(report['totals']['bias_pct'], 0)
        self.assertEqual(report['totals']['within_tolerance_pct'], 0)
        self.assertEqual(report['questions']['mae_pct'], 10)
        self.assertEqual(self.db.attempts.find_one()['marks'], 76)
        self.assertEqual(self.db.evaluations.find_one()['marks'], 30)
        self.assertEqual(self.db.attempts.count_documents({}), 1)

    def test_manual_mismatched_and_unmatched_results_excluded(self):
        self.db.attempts.update_one({}, {'$set': {'source': 'manual'}})
        report = self.preview().get_json()['report']
        self.assertEqual(report['statuses']['manual_result'], 2)
        self.assertEqual(report['totals']['compared'], 1)
        response = self.preview('source,result_id,teacher_marks,max_marks\nworkspace,missing,1,10\nlegacy,' + str(self.legacy) + ',20,100\n')
        report = response.get_json()['report']
        self.assertEqual(report['statuses'], {'unmatched': 1, 'max_marks_mismatch': 1})
        self.assertIsNone(report['totals']['mae_pct'])

    def test_snapshot_save_and_latest_keep_prior_scores(self):
        report = self.preview().get_json()['report']
        self.assertEqual(self.client.post(f"{self.base}/{report['_id']}/save", headers=self.headers).status_code, 200)
        self.db.attempts.update_one({}, {'$set': {'marks': 72}})
        latest = self.client.post(f"{self.base}/{report['_id']}/compare-latest", headers=self.headers).get_json()['report']
        self.assertEqual(latest['rows'][0]['ai_marks'], 72)
        prior = self.client.get(f"{self.base}/{report['_id']}", headers=self.headers).get_json()['report']
        self.assertEqual(prior['rows'][0]['ai_marks'], 76)
        self.assertEqual(latest['parent_id'], prior['_id'])

    def test_permissions_and_auth(self):
        self.assertEqual(self.client.get(self.base).status_code, 401)
        report = self.preview().get_json()['report']
        headers = {**self.headers, 'X-User': 'other@example.com'}
        for suffix in ('', '/export'):
            self.assertEqual(self.client.get(f"{self.base}/{report['_id']}{suffix}", headers=headers).status_code, 403)
        self.assertEqual(self.client.get(self.base + '/template?course_code=FS3401', headers=headers).status_code, 403)

    def test_bad_numbers_duplicates_and_formulas_rejected(self):
        for marks in ('NaN', 'Infinity', '-1', '101', '=1+1'):
            self.assertEqual(self.preview(f'source,result_id,teacher_marks,max_marks\nworkspace,sheet,{marks},100\n').status_code, 400)
        self.assertEqual(self.preview('source,result_id,teacher_marks,max_marks\nworkspace,sheet,70,100\nworkspace,sheet,71,100\n').status_code, 400)

    def test_natural_matching_ambiguous_not_guessed_and_ids_checked(self):
        self.db.sheets.insert_one({'_id': 'duplicate', 'course_code': 'FS3401', 'exam_id': 'October', 'roll_number': '001'})
        report = self.preview('roll_number,exam_id,teacher_marks,max_marks\n001,October,72,100\n').get_json()['report']
        self.assertEqual(report['rows'][0]['status'], 'ambiguous')
        report = self.preview('source,result_id,roll_number,teacher_marks,max_marks\nworkspace,sheet,WRONG,72,100\n').get_json()['report']
        self.assertEqual(report['rows'][0]['status'], 'identity_mismatch')

    def test_template_omits_ai_marks_and_xlsx_roundtrip(self):
        response = self.client.get(self.base + '/template?course_code=FS3401', headers=self.headers)
        wb = load_workbook(io.BytesIO(response.data))
        ws = wb.active
        self.assertEqual(ws['E2'].value, '001')
        self.assertIsNone(ws['G2'].value)
        ws['G2'] = 72
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        labels = parse_upload(FileStorage(stream=output, filename='marks.xlsx'), 'FS3401')
        self.assertEqual(len(labels), 1)
        self.assertEqual(labels[0]['roll_number'], '001')

    def test_calibration_and_test_papers_cannot_overlap(self):
        report = self.preview().get_json()['report']
        held_out = self.preview(split='test').get_json()['report']
        self.client.post(f"{self.base}/{report['_id']}/save", headers=self.headers)
        self.assertEqual(self.preview(split='test').status_code, 400)
        self.assertEqual(self.client.post(f"{self.base}/{held_out['_id']}/save", headers=self.headers).status_code, 400)

    def test_export_treats_feedback_as_text(self):
        report = self.preview('source,result_id,teacher_marks,max_marks,teacher_feedback\nworkspace,sheet,72,100,=1+1\n').get_json()['report']
        response = self.client.get(f"{self.base}/{report['_id']}/export", headers=self.headers)
        wb = load_workbook(io.BytesIO(response.data))
        ws = wb['Teacher marks']
        self.assertEqual(wb.active.title, 'Comparison summary')
        self.assertEqual(ws['I2'].value, '=1+1')
        self.assertEqual(ws['I2'].data_type, 's')

    def test_malformed_workbook_is_a_validation_error(self):
        response = self.client.post(self.base + '/preview', headers=self.headers, data={
            'course_code': 'FS3401', 'name': 'Invalid workbook',
            'file': (io.BytesIO(b'not an Excel workbook'), 'marks.xlsx')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.db.teacher_comparisons.count_documents({}), 0)

    def test_missing_difficult_paper_does_not_claim_improvement(self):
        baseline = self.preview(f'source,result_id,teacher_marks,max_marks\nworkspace,sheet,72,100\nlegacy,{self.legacy},50,50\n').get_json()['report']
        self.client.post(f"{self.base}/{baseline['_id']}/save", headers=self.headers)
        self.db.evaluations.delete_one({'_id': self.legacy})
        latest = self.client.post(f"{self.base}/{baseline['_id']}/compare-latest", headers=self.headers).get_json()['report']
        self.assertLess(latest['totals']['mae_pct'], baseline['totals']['mae_pct'])
        self.assertEqual(latest['change']['totals']['paired'], 1)
        self.assertEqual(latest['change']['totals']['mae_change_pct'], 0)
        self.assertEqual(latest['change']['totals']['no_longer_comparable'], 1)

    def test_paired_improvement_survives_saving_and_reopening(self):
        baseline = self.preview('source,result_id,teacher_marks,max_marks\nworkspace,sheet,72,100\n').get_json()['report']
        self.client.post(f"{self.base}/{baseline['_id']}/save", headers=self.headers)
        attempt = self.db.attempts.find_one({'_id': 'attempt'})
        attempt.update(_id='new-attempt', marks=73)
        self.db.attempts.insert_one(attempt)
        self.db.sheets.update_one({'_id': 'sheet'}, {'$set': {'active_attempt': 'new-attempt'}})
        latest = self.client.post(f"{self.base}/{baseline['_id']}/compare-latest", headers=self.headers).get_json()['report']
        self.assertEqual(latest['change']['totals']['mae_change_pct'], -3)
        self.assertEqual(latest['change']['totals']['improved'], 1)
        self.assertEqual(latest['change']['totals']['same_prediction'], 0)
        self.client.post(f"{self.base}/{latest['_id']}/save", headers=self.headers)
        reopened = self.client.get(f"{self.base}/{latest['_id']}", headers=self.headers).get_json()['report']
        self.assertEqual(reopened['change'], latest['change'])

    def test_no_pairs_returns_no_accuracy_claim(self):
        baseline = self.preview('source,result_id,teacher_marks,max_marks\nworkspace,sheet,72,100\n').get_json()['report']
        self.client.post(f"{self.base}/{baseline['_id']}/save", headers=self.headers)
        self.db.sheets.delete_one({'_id': 'sheet'})
        latest = self.client.post(f"{self.base}/{baseline['_id']}/compare-latest", headers=self.headers).get_json()['report']
        self.assertEqual(latest['change']['totals']['paired'], 0)
        self.assertIsNone(latest['change']['totals']['mae_change_pct'])

    def test_relative_error_matches_teacher_baseline_and_handles_zero(self):
        self.db.attempts.update_one({}, {'$set': {'marks': 54}})
        report = self.preview(f'source,result_id,teacher_marks,max_marks\nworkspace,sheet,50,100\nlegacy,{self.legacy},0,50\n').get_json()['report']
        self.assertEqual(report['totals']['mape_pct'], 8)
        self.assertEqual(report['totals']['relative_compared'], 1)
        self.assertEqual(report['totals']['zero_teacher_rows'], 1)
        self.assertEqual(report['totals']['compared'], 2)
        self.client.post(f"{self.base}/{report['_id']}/save", headers=self.headers)
        self.db.attempts.update_one({}, {'$set': {'marks': 52}})
        latest = self.client.post(f"{self.base}/{report['_id']}/compare-latest", headers=self.headers).get_json()['report']
        self.assertEqual(latest['change']['totals']['before_mape_pct'], 8)
        self.assertEqual(latest['change']['totals']['after_mape_pct'], 4)


if __name__ == '__main__':
    unittest.main()

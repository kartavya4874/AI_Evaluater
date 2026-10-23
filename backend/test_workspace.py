import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
from concurrent.futures import ThreadPoolExecutor

import mongomock
from openpyxl import load_workbook
from utils.workspace import Workspace, safe_path, excel_bytes
from utils.result_validation import validate_result


def grade(marks=65):
    return {'marks_awarded': marks, 'feedback': 'Valid feedback', 'strengths': [], 'missing_points': []}


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        course = root / 'FS3401_MSE'
        course.mkdir()
        (course / 'model_answer.pdf').write_bytes(b'answer key')
        (course / '101.pdf').write_bytes(b'student')
        self.db = mongomock.MongoClient().db
        self.service = Workspace(self.db)
        self.sheet_id = self.service.prepare(str(root), 'exam-1', 'mst', 100, 'admin')[0]
        self.pdf = Mock()
        self.pdf.extract_text_from_pdf.return_value = 'Extracted complete text'
        self.llm = Mock()
        self.llm.evaluate_answer.return_value = grade()

    def process(self, force=False):
        self.service.enqueue(self.sheet_id, 'admin', force)
        job = self.service.claim('worker')
        self.assertIsNotNone(job)
        self.service.process(job, self.pdf, self.llm, 'worker')
        return self.service.snapshot()[0]

    def test_forced_run_refreshes_scores_but_reuses_preprocessing(self):
        observed = []
        def evaluate(*args, **kwargs):
            cache = self.llm.structured_cache
            observed.append((cache('Extract rubric'), cache('Grade ONLY question 1')))
            cache('Extract rubric', {'questions': []})
            cache('Grade ONLY question 1', {'marks_awarded': 65})
            return grade()
        self.llm.evaluate_answer.side_effect = evaluate
        self.process()
        result = self.process(force=True)
        self.assertEqual(observed, [(None, None), ({'questions': []}, None)])
        self.assertEqual(self.pdf.extract_text_from_pdf.call_count, 2)
        attempt = self.db.attempts.find_one({'_id': result['active']['_id']})
        metrics = attempt['performance']
        self.assertFalse(metrics['result_cache_hit'])
        self.assertGreaterEqual(metrics['cache_hits'], 3)
        self.assertGreaterEqual(metrics['processing_seconds'], metrics['grading_seconds'])

    def test_zero_is_valid(self):
        self.llm.evaluate_answer.return_value = grade(0)
        sheet = self.process()
        self.assertEqual(sheet['active']['marks'], 0)
        self.assertEqual(sheet['active']['status'], 'completed')

    def test_ocr_progress_counts_completed_pages(self):
        self.pdf.extract_text_from_pdf.side_effect = ValueError('scanned PDF')
        self.pdf.convert_pdf_to_images.return_value = [Mock(), Mock()]
        observed = []
        def ocr(*args, **kwargs):
            progress = self.db.jobs.find_one({'_id': self.sheet_id})['progress']
            observed.append((progress['completed'], progress['total']))
            return 'Read page'
        self.pdf.extract_text_from_images_via_gemini.side_effect = ocr
        sheet = self.process()
        self.assertEqual(observed, [(0, 2), (1, 2), (0, 2), (1, 2)])
        self.assertIn('extracted_student_page_2_of_2', [e['stage'] for e in sheet['job']['events']])
        self.assertIsNone(sheet['job']['progress'])

    def test_activity_events_persist_for_success_and_failure(self):
        sheet = self.process()
        stages = [event['stage'] for event in sheet['job']['events']]
        self.assertEqual(stages[0], 'queued')
        self.assertIn('grading', stages)
        self.assertEqual(stages[-1], 'completed')
        self.assertTrue(all(event.get('at') for event in sheet['job']['events']))
        self.llm.evaluate_answer.side_effect = ValueError('Invalid source')
        failed = self.process(force=True)
        self.assertEqual(failed['job']['events'][-1]['stage'], 'error')
        self.assertEqual(failed['job']['events'][-1]['message'], 'Invalid source')
        self.assertIsNone(failed['job']['progress'])

    def test_failure_has_no_marks_or_grade_cache(self):
        self.llm.evaluate_answer.side_effect = ValueError('Invalid source')
        sheet = self.process()
        self.assertIsNone(sheet['active'])
        self.assertIsNone(sheet['attempts'][0]['marks'])
        self.assertEqual(sheet['job']['state'], 'error')
        self.assertEqual(self.db.workspace_cache.count_documents({'result': {'$exists': True}}), 0)

    def test_retry_preserves_history_and_active_result(self):
        first = self.process()['active']['_id']
        self.llm.evaluate_answer.side_effect = ValueError('Invalid source')
        failed = self.process(force=True)
        self.assertEqual(failed['active']['_id'], first)
        self.assertEqual(len(failed['attempts']), 2)
        self.llm.evaluate_answer.side_effect = None
        self.llm.evaluate_answer.return_value = grade(70)
        final = self.process(force=True)
        self.assertEqual(final['active']['marks'], 70)
        self.assertEqual(len(final['attempts']), 3)

    def test_duplicate_enqueue_is_idempotent(self):
        one = self.service.enqueue(self.sheet_id, 'admin')
        two = self.service.enqueue(self.sheet_id, 'admin', True)
        self.assertEqual(one['attempt_id'], two['attempt_id'])

    def test_retry_claim_clears_previous_error(self):
        self.service.enqueue(self.sheet_id, 'admin')
        self.db.jobs.update_one({'_id': self.sheet_id}, {'$set': {'error': 'Old upload timeout'}})
        job = self.service.claim('worker')
        self.assertEqual(job['state'], 'processing')
        self.assertIsNone(job['error'])

    def test_concurrent_enqueue_has_one_job(self):
        with ThreadPoolExecutor(max_workers=8) as executor:
            jobs = list(executor.map(lambda _: self.service.enqueue(self.sheet_id, 'admin', True), range(20)))
        self.assertEqual(len({j['attempt_id'] for j in jobs}), 1)
        self.assertEqual(self.db.jobs.count_documents({}), 1)

    def test_changed_inputs_invalidate_active_result(self):
        self.process()
        (Path(self.temp.name) / 'FS3401_MSE/model_answer.pdf').write_bytes(b'changed key')
        self.service.prepare(self.temp.name, 'exam-1', 'mst', 100, 'admin')
        sheet = self.service.snapshot()[0]
        self.assertIsNone(sheet['active'])
        self.assertEqual(len(sheet['attempts']), 1)
        self.assertEqual(self.service.enqueue(self.sheet_id, 'admin')['state'], 'queued')

    def test_pinned_result_survives_successful_retry(self):
        original = self.process()['active']['_id']
        self.db.sheets.update_one({'_id': self.sheet_id}, {'$set': {'active_pinned': True}})
        self.llm.evaluate_answer.return_value = grade(75)
        sheet = self.process(force=True)
        self.assertEqual(sheet['active']['_id'], original)
        self.assertEqual(len(sheet['attempts']), 2)

    def test_cancelled_job_does_not_grade(self):
        self.service.enqueue(self.sheet_id, 'admin')
        job = self.service.claim('worker')
        self.db.jobs.update_one({'_id': self.sheet_id}, {'$set': {'cancel_requested': True}})
        self.service.process(job, self.pdf, self.llm, 'worker')
        self.llm.evaluate_answer.assert_not_called()
        self.assertEqual(self.db.jobs.find_one() ['state'], 'cancelled')

    def test_stop_during_grading_prevents_further_subrequests_and_result(self):
        def evaluate(*args, **kwargs):
            self.db.jobs.update_one({'_id': self.sheet_id}, {'$set': {'cancel_requested': True}})
            self.llm.check_cancelled()
            self.fail('Grading continued after cancellation')
        self.llm.evaluate_answer.side_effect = evaluate
        sheet = self.process()
        self.assertEqual(sheet['job']['state'], 'cancelled')
        self.assertIsNone(sheet['active'])
        self.assertEqual(self.db.attempts.find_one()['status'], 'cancelled')

    def test_expired_worker_cannot_publish(self):
        from utils.workspace import now
        self.db.worker_leases.insert_one({'_id': 'inference', 'owner': 'different', 'expires': now()})
        sheet = self.process()
        self.assertIsNone(sheet['active'])
        self.llm.evaluate_answer.assert_not_called()

    def test_management_guard_blocks_enqueue(self):
        with self.service.reserve([self.sheet_id]):
            self.assertEqual(self.service.enqueue(self.sheet_id, 'admin')['state'], 'managing')
        self.assertEqual(self.db.jobs.count_documents({}), 0)

    def test_crash_after_result_save_is_reconciled(self):
        self.service.enqueue(self.sheet_id, 'admin')
        job = self.service.claim('old-worker')
        with patch.object(self.db.sheets, 'update_one', side_effect=RuntimeError('database write interrupted')):
            with self.assertRaises(RuntimeError):
                self.service.process(job, self.pdf, self.llm, 'old-worker')
        self.assertEqual(self.db.attempts.find_one()['status'], 'completed')
        self.service.recover_interrupted('new-worker')
        sheet = self.service.snapshot()[0]
        self.assertEqual(sheet['active']['marks'], 65)
        self.assertEqual(sheet['job']['state'], 'completed')

    def test_interrupted_attempt_is_retryable(self):
        self.service.enqueue(self.sheet_id, 'admin')
        job = self.service.claim('old-worker')
        self.db.attempts.insert_one({'_id': job['attempt_id'], 'sheet_id': self.sheet_id, 'status': 'processing'})
        self.service.recover_interrupted('new-worker')
        self.assertEqual(self.db.attempts.find_one()['status'], 'processing')
        self.assertEqual(self.service.enqueue(self.sheet_id, 'admin')['state'], 'queued')

    def test_remove_sheet_deletes_only_selected_scope(self):
        self.process()
        self.service.manage([self.sheet_id], 'sheet', 'admin')
        self.assertEqual(self.db.sheets.count_documents({}), 0)
        self.assertEqual(self.db.jobs.count_documents({}), 0)
        self.assertEqual(self.db.attempts.count_documents({}), 0)
        self.assertEqual(self.db.workspace_cache.count_documents({}), 0)
        self.assertTrue((Path(self.temp.name) / 'FS3401_MSE/101.pdf').exists())

    def test_page_retry_reuses_successful_ocr_pages(self):
        self.pdf.extract_text_from_pdf.side_effect = ValueError('scanned')
        self.pdf.convert_pdf_to_images.side_effect = lambda path: [Mock(), Mock()]
        self.pdf.extract_text_from_images_via_gemini.side_effect = ['page one', RuntimeError('page two failed')]
        self.assertEqual(self.process()['job']['state'], 'error')
        self.pdf.extract_text_from_images_via_gemini.reset_mock()
        self.pdf.extract_text_from_images_via_gemini.side_effect = None
        self.pdf.extract_text_from_images_via_gemini.return_value = 'extracted page'
        self.assertEqual(self.process()['job']['state'], 'completed')
        # Retry needs only key page two and the two student pages.
        self.assertEqual(self.pdf.extract_text_from_images_via_gemini.call_count, 3)

    def test_new_exam_and_changed_key_do_not_reuse(self):
        root = self.temp.name
        second = self.service.prepare(root, 'exam-2', 'mst', 100, 'admin')[0]
        self.assertNotEqual(second, self.sheet_id)
        (Path(root) / 'FS3401_MSE/model_answer.pdf').write_bytes(b'new key')
        third = self.service.prepare(root, 'exam-1', 'mst', 100, 'admin')[0]
        self.assertEqual(third, self.sheet_id)
        self.assertEqual(self.db.sheets.count_documents({}), 2)

    def test_reupload_reuses_same_sheet(self):
        ids = self.service.prepare(self.temp.name, 'exam-1', 'mst', 100, 'admin')
        self.assertEqual(ids, [self.sheet_id])

    def test_cache_deletion_preserves_results(self):
        self.process()
        self.service.manage([self.sheet_id], 'cache', 'admin')
        self.assertEqual(self.db.workspace_cache.count_documents({}), 0)
        self.assertIsNotNone(self.service.snapshot()[0]['active'])

    def test_record_deletion_allows_new_attempt(self):
        self.process()
        self.service.manage([self.sheet_id], 'both', 'admin')
        self.assertEqual(self.db.attempts.count_documents({}), 0)
        self.assertIsNone(self.service.snapshot()[0]['active'])
        self.assertEqual(self.process()['active']['marks'], 65)

    def test_management_rejects_active_jobs(self):
        self.service.enqueue(self.sheet_id, 'admin')
        with self.assertRaises(ValueError):
            self.service.manage([self.sheet_id], 'both', 'admin')

    def test_ocr_failure_stops_grading(self):
        self.pdf.extract_text_from_pdf.side_effect = ValueError('scanned')
        self.pdf.convert_pdf_to_images.side_effect = RuntimeError('bad PDF')
        sheet = self.process()
        self.assertEqual(sheet['job']['state'], 'error')
        self.llm.evaluate_answer.assert_not_called()

    def test_export_excludes_feedback_and_ocr_and_formulas(self):
        self.process()
        rows = self.service.snapshot()
        rows[0]['roll_number'] = '=1+1'
        rows[0]['active']['verified'] = True
        workbook = load_workbook(excel_bytes(rows))
        ws = workbook.active
        values = list(ws.values)
        self.assertEqual(values[1][4], 65)
        self.assertNotIn('Valid feedback', str(values))
        self.assertNotIn('Extracted complete text', str(values))
        self.assertEqual(ws['C2'].data_type, 's')

    def test_invalid_results_rejected(self):
        for marks in (-1, 101, float('nan'), True, '65'):
            with self.assertRaises(ValueError):
                validate_result(grade(marks), 100)
        with self.assertRaises(ValueError):
            validate_result({**grade(0), 'needs_review': True}, 100)

    def test_traversal_rejected(self):
        for path in ('../outside.pdf', '/etc/file.pdf', 'C:/outside.pdf', 'folder/../../file.pdf'):
            with self.assertRaises(ValueError):
                safe_path(self.temp.name, path)


if __name__ == '__main__':
    unittest.main()
